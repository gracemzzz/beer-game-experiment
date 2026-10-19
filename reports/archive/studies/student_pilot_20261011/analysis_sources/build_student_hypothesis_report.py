"""Read-only, versioned H1/H2 analysis of one saved student-protocol campaign.

No model client, key prompt, experiment mutation, or paid calls. Use --run PATH.
New analysis sources and input hashes determine a separate analysis fingerprint.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))

import numpy as np
import pandas as pd

from beergame_study.student_hypotheses import analyze_run

ROLES = ('retailer', 'wholesaler', 'distributor', 'factory')
VERSION = 'student-hypothesis-report-v1'
WORD_REPORT = 'Student_Pilot_Report.docx'
MARKDOWN_REPORT = 'REPORT.md'
WORKBOOK = 'DETAILS.xlsx'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def plain(value):
    if isinstance(value, (np.integer, np.floating)):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, dict):
        return {str(k): plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [plain(v) for v in value]
    if isinstance(value, Path):
        return str(value)
    return value


def show(value):
    if value is None or (not isinstance(value, (list, dict)) and pd.isna(value)):
        return '—'
    if isinstance(value, (bool, np.bool_)):
        return 'yes' if value else 'no'
    if isinstance(value, (float, np.floating)):
        if value == 0:
            return '0'
        if abs(value) < .001:
            return f'{value:.3g}'
        return f'{value:,.6g}'
    return str(value).replace('|', '/').replace('\n', ' ')


def supplementary_tables(run):
    rounds = pd.read_csv(run / 'rounds.csv')
    try:
        tail = pd.read_csv(run / 'continuation.csv')
    except (FileNotFoundError, pd.errors.EmptyDataError):
        tail = pd.DataFrame(columns=rounds.columns)
    agent = rounds[rounds.phase == 'agent']
    cost_rows, bullwhip_rows, agreement_rows, behavior_rows = [], [], [], []
    for (gid, condition, block), rows in rounds.groupby(['group_id', 'condition', 'block']):
        continuation = tail[tail.group_id == gid]
        for role, r in rows.groupby('role'):
            choices = r[r.phase == 'agent']
            t = continuation[continuation.role == role].sort_values('week')
            behavior_rows.append(dict(group_id=gid, condition=condition, role=role,
                valid_choices=int(choices.valid_decision.sum()),
                zero_orders=int(((choices.order == 0) & choices.valid_decision).sum()),
                max_order=float(choices.order.max()) if len(choices) else np.nan,
                concerning_zeros=int(((choices.order == 0) & (choices.inventory == 0) &
                    (choices.supply_line == 0) & (choices.backlog > 0) & choices.valid_decision).sum()),
                backlog_weeks=int((r.backlog > 0).sum()), max_backlog=int(r.backlog.max()),
                observed_holding_cost=float(.5 * r.inventory.sum()), observed_backlog_cost=float(r.backlog.sum()),
                cumulative_shipped_over_demand=float(r.shipped.sum() / r.demand.sum()) if r.demand.sum() else np.nan,
                terminal_inventory=float(t.inventory.iloc[-1]) if len(t) else np.nan,
                terminal_backlog=float(t.backlog.iloc[-1]) if len(t) else np.nan,
                terminal_preorder_supply_line=float(t.supply_line.iloc[-1]) if len(t) else np.nan))
        for label, part in (('familiarization', rows[rows.phase == 'warmup']),
                            ('model-choice weeks', rows[rows.phase == 'agent']),
                            ('automatic continuation', continuation)):
            cost_rows.append(dict(group_id=gid, condition=condition, block=block,
                                  scoring_phase=label, holding_cost=float(.5 * part.inventory.sum()),
                                  backlog_cost=float(part.backlog.sum()), total_cost=float(part.cost.sum())))
        config = json.loads((run / 'manifest.json').read_text(encoding='utf-8'))['config']
        for label, first in (('full game including familiarization', 1),
                             ('model-choice weeks', config['warmup_weeks'] + 1), ('late-window sensitivity', 9)):
            part = rows[rows.week >= first]
            if part.empty:
                continue
            customer = part[part.role == 'retailer'].sort_values('week').demand.to_numpy()
            customer_var = float(np.var(customer, ddof=1)) if len(customer) > 1 else np.nan
            for role, r in part.groupby('role'):
                order_var, demand_var = float(r.order.var(ddof=1)), float(r.demand.var(ddof=1))
                bullwhip_rows.append(dict(group_id=gid, block=block, condition=condition, role=role,
                    window=label, first_week=first, last_week=int(r.week.max()), n_weeks=len(r),
                    order_variance=order_var, local_demand_variance=demand_var,
                    customer_demand_variance=customer_var,
                    bullwhip_vs_local=order_var / demand_var if demand_var > 1e-12 else np.nan,
                    bullwhip_vs_customer=order_var / customer_var if customer_var > 1e-12 else np.nan))
    for block, rows in agent.groupby('block'):
        a = rows[rows.condition == 'short'].set_index(['week', 'role'])
        b = rows[rows.condition == 'adequate'].set_index(['week', 'role'])
        joined = a[['order', 'demand', 'valid_decision', 'fallback']].join(
            b[['order', 'demand', 'valid_decision', 'fallback']], how='inner', lsuffix='_short', rsuffix='_adequate')
        clean = joined.valid_decision_short & joined.valid_decision_adequate & ~joined.fallback_short & ~joined.fallback_adequate
        j = joined[clean]
        agreement_rows.append(dict(block=int(block), matched_valid_week_role_choices=len(j),
            identical_orders=int((j.order_short == j.order_adequate).sum()),
            agreement_fraction=float((j.order_short == j.order_adequate).mean()) if len(j) else np.nan,
            mean_absolute_order_difference=float((j.order_short - j.order_adequate).abs().mean()) if len(j) else np.nan,
            short_orders_equal_local_demand=bool((j.order_short == j.demand_short).all()) if len(j) else False,
            adequate_orders_equal_local_demand=bool((j.order_adequate == j.demand_adequate).all()) if len(j) else False))
    calibration = json.loads((run / 'calibration_recommendation.json').read_text(encoding='utf-8'))
    return {'cost_components': pd.DataFrame(cost_rows), 'bullwhip_windows': pd.DataFrame(bullwhip_rows),
            'behavior_checks': pd.DataFrame(behavior_rows),
            'order_agreement': pd.DataFrame(agreement_rows), 'calibration_checks': pd.DataFrame(calibration['screen_rows'])}


def plots(run, output):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    frame = pd.read_csv(run / 'rounds.csv')
    fig, axes = plt.subplots(2, 2, figsize=(11, 6), sharex=True)
    for ax, role in zip(axes.flat, ROLES):
        for (condition, block), rows in frame[frame.role == role].groupby(['condition', 'block']):
            ax.plot(rows.week, rows.order, '--' if condition == 'adequate' else '-',
                    label=f'{condition}, block {block}', linewidth=2, alpha=.8)
            ax.plot(rows.week, rows.demand, ':', linewidth=1,
                    label=f'local demand: {condition}, block {block}', alpha=.8)
        ax.set_title(role.title()); ax.set_ylabel('Cases'); ax.legend(fontsize=7)
    axes[1, 0].set_xlabel('Game week'); axes[1, 1].set_xlabel('Game week')
    fig.suptitle('Descriptive order and demand trajectories; overlapping lines are retained')
    fig.tight_layout(); fig.savefig(output / 'order_demand_checks.png', dpi=160); plt.close(fig)
    measured = frame[(frame.phase == 'agent') & frame.usage_observed]
    fig, ax = plt.subplots(figsize=(9, 4))
    for condition, values in measured.groupby('condition'):
        values = np.sort(values.output_tokens.to_numpy())
        ax.step(values, np.arange(1, len(values) + 1) / len(values), where='post', label=condition)
    ax.set_xlabel('Actual generated tokens, including reasoning'); ax.set_ylabel('Fraction of measured calls')
    ax.set_title('Token-use distributions: repeated decisions, not independent students')
    if len(measured):
        ax.legend()
    else:
        ax.text(.5, .5, 'No measured API token usage in this run', ha='center', transform=ax.transAxes)
    fig.tight_layout(); fig.savefig(output / 'actual_token_ecdf.png', dpi=160); plt.close(fig)


def resolve_report_run(report):
    """Find source evidence from either a published report or a run report."""
    if report is None:
        return None
    path = Path(report).resolve()
    directory = path.parent if path.is_file() else path
    if (directory / 'manifest.json').is_file():
        return directory
    provenance = directory / 'analysis_manifest.json'
    if provenance.is_file():
        saved = json.loads(provenance.read_text(encoding='utf-8'))
        return Path(saved.get('run_directory', ROOT / 'data/runs' / saved['run_id']))
    return None


def build_report(directory, *, update_pointer=True, reports_root=None):
    run = Path(directory).resolve()
    result = analyze_run(run)
    tables = {k: v for k, v in result.items() if isinstance(v, pd.DataFrame)}
    tables.update(supplementary_tables(run))
    manifest = json.loads((run / 'manifest.json').read_text(encoding='utf-8'))
    selection = json.loads((run / 'selection.json').read_text(encoding='utf-8'))
    config = manifest['config']
    saved_summary = json.loads((run / 'report_summary.json').read_text(encoding='utf-8'))
    progress = json.loads((run / 'progress.json').read_text(encoding='utf-8'))
    completed_pilot = run.name == 'student_api_b609eca42866feb0' and progress['phase'] == 'complete'
    analysis_date = datetime.now().astimezone().strftime('%d %B %Y')
    inputs = ['manifest.json', 'selection.json', 'progress.json', 'rounds.csv', 'continuation.csv',
              'calibration.csv', 'calibration_recommendation.json', 'group_outcomes.csv',
              'beta_fits.csv', 'role_metrics.csv', 'report_summary.json', 'rules_checks.json']
    inputs += [f for f in ('independent_audit.json', 'calibration_import.json', 'calibration_reuse.json', 'report_annotation.json') if (run / f).exists()]
    inputs += [str(p.relative_to(run)).replace('\\', '/') for p in sorted((run / 'requests').glob('*.json'))]
    inputs = [name for name in inputs if (run / name).exists()]
    input_hashes = {name: sha(run / name) for name in inputs}
    sources = {str(p.relative_to(ROOT)).replace('\\', '/'): sha(p) for p in
               (Path(__file__), ROOT / 'src/beergame_study/student_hypotheses.py')}
    identity = dict(analysis_version=VERSION, run_protocol=manifest['protocol'], run_id=run.name,
                    input_sha256=input_hashes, analysis_source_sha256=sources)
    reference_names = ['data/source/ThoughtsReflection_Paragraph.docx', 'docs/PROTOCOL.md',
                       'docs/RESEARCH_DESIGN_20261011.md', 'src/beergame_study/analysis.py']
    identity['reference_source_sha256'] = {name: sha(ROOT / name) for name in reference_names if (ROOT / name).exists()}
    calibration_source = Path(result['summary']['calibration_source'])
    calibration_inputs = [calibration_source / name for name in ('manifest.json', 'progress.json',
        'calibration.csv', 'calibration_recommendation.json', 'calibration_import.json') if (calibration_source / name).exists()]
    calibration_inputs += sorted((calibration_source / 'requests').glob('*.json'))
    calibration_inputs += sorted((calibration_source / 'source_implementation').glob('*.py'))
    identity['calibration_evidence'] = dict(directory=str(calibration_source),
        file_sha256={str(p.relative_to(calibration_source)): sha(p) for p in calibration_inputs})
    fingerprint = hashlib.sha256(json.dumps(identity, sort_keys=True).encode()).hexdigest()
    reports_root = Path(reports_root or ROOT / 'reports').resolve()
    folder = 'student_pilot_20261011' if run.name == 'student_api_b609eca42866feb0' else run.name
    output = reports_root / folder
    previous_manifest = output / 'analysis_manifest.json'
    if previous_manifest.exists():
        previous = json.loads(previous_manifest.read_text(encoding='utf-8'))
        if previous['analysis_fingerprint'] != fingerprint:
            archived = reports_root / 'archive/analysis_versions' / run.name / previous['analysis_fingerprint'][:16]
            # Both resolved targets must stay within the reports directory before moving.
            output.resolve().relative_to(reports_root)
            archived.resolve().relative_to(reports_root)
            if archived.exists():
                raise FileExistsError(f'An archived report already occupies {archived}')
            archived.parent.mkdir(parents=True, exist_ok=True)
            output.rename(archived)
    elif output.exists() and any(output.iterdir()):
        raise FileExistsError(f'Report directory has no analysis manifest: {output}')
    output.mkdir(parents=True, exist_ok=True)
    for name in sources:
        destination = output / 'analysis_sources' / Path(name).name
        destination.parent.mkdir(exist_ok=True)
        destination.write_bytes((ROOT / name).read_bytes())
    for name in identity['reference_source_sha256']:
        destination = output / 'reference_sources' / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes((ROOT / name).read_bytes())
    sections = []

    def text(title, paragraphs):
        sections.append(('text', title, paragraphs))

    def table(title, key, columns=None):
        frame = tables[key]
        if columns:
            frame = frame[[c for c in columns if c in frame.columns]]
        sections.append(('table', title, frame))

    text('Conclusion and scope', [
        f'This is the full H1/H2 analysis of saved run {run.name}, produced without additional model calls. '
        'It adds the hypothesis checks omitted from the earlier outcome report. It is a post-hoc analysis document; '
        'it does not revise the saved experiment or claim prospective registration of this new analysis code.',
        f'The frozen plan selected {selection["groups_per_condition"]} chain pair(s), with '
        f'{selection["short_tokens"]:,} versus {selection["adequate_tokens"]:,} generated-token caps. '
        f'The calibration manipulation screen passed: {selection["calibration_passed"]}. '
        'Read the separate H1 and H2 eligibility decisions below; withheld tests are missing, rather than zero or a fabricated p-value.',
        'For the completed 11 October pilot, both chains scored 156 and matched on every valid week-role order. '
        'H1 cannot be estimated from usable coefficients; H2 has a descriptive paired difference of zero. '
        'Neither hypothesis has a valid confirmatory conclusion from this pilot. Identical observations do not establish equivalence or the absence of an effect.',
        'The evidence consists of configured LLM simulations. It contains no real students, and weekly calls are repeated interacting decisions. '
        'The analysis unit is an independently initialized four-role supply chain, paired by block across conditions.'])
    # The specific pilot conclusion is clearly isolated when the same exporter is used on another run.
    if not completed_pilot:
        sections[0][2][2] = 'This report analyzes the selected run; numerical outcomes and individual eligibility decisions are in the tables below. No withheld test establishes equivalence or absence of an effect.'
    if config['mode'] == 'offline':
        sections[0][2].insert(0, 'SYNTHETIC SOFTWARE BENCHMARK: this run contains no live GPT decisions. '
            'The configured model and high-effort settings name the intended API protocol, not computation performed by this benchmark. '
            'Synthetic trajectories cannot validate the persona, manipulation or behavioral hypotheses.')
    text('Original hypotheses, estimands and planned tests', [
        'Original H1: "Within an information-sharing environment, a binding decision-time constraint increases the underweighting of outstanding replenishment orders relative to adequate decision time." '
        'The repo operationalizes this as a lower supply-line coefficient beta under the short allowance. '
        'For block b, D_beta,b = mean(beta_short,b,r over all four roles) - mean(beta_adequate,b,r over all four roles). '
        'The null is no paired condition difference; the predicted direction is E[D_beta] < 0. A chain mean exists only if every role fit is usable.',
        'Original H2: "Within an information-sharing environment, supply chains operating under binding decision-time constraints incur higher total inventory-holding and backlog costs than supply chains operating with adequate decision time." '
        'D_cost,b = C_short,b - C_adequate,b, with C = sum over roles and all scored weeks of (0.50 * inventory + 1.00 * backlog). '
        'The null is no paired condition difference; the predicted direction is E[D_cost] > 0. Costs include fixed familiarization and automatic continuation. '
        'H2 does not require identifiable beta and is checked independently.',
        'Following the earlier workflow, eligible comparisons use a two-sided paired sign-flip test on mean chain-block differences. '
        'For n <= 16, enumerate all 2^n sign vectors and count abs(mean(s_b * D_b)) >= abs(mean(D_b)). '
        'For larger n, use the fixed-seed 50,000-draw Monte Carlo variant with a plus-one correction. '
        'This requires sign exchangeability under the null; a mean-equality statement alone does not establish that assumption. '
        'Predicted direction is reported separately from a two-sided p-value.',
        'Holm adjustment retains the planned family of two hypotheses. If one test is unavailable, a p=1 placeholder is used only internally for the other adjustment; '
        'the unavailable p-value stays blank in reports. The inherited eight-pair minimum is a software inference floor, not a power calculation. '
        'This new analysis additionally withholds inferential results when the recorded token-manipulation screen failed. '
        'No independent-pair confidence interval is manufactured from one pair; no bootstrap over weekly decisions is used.',
        'The hypotheses originate in data/source/ThoughtsReflection_Paragraph.docx and the legacy docs/PROTOCOL.md. '
        'Their mapping is retained while the active prompt/model/high-effort/single-call token protocol differs from the historical none-effort, two-call implementation. '
        'This is an exploratory analysis workflow, not a registered human experiment.'])
    tables['design_mapping'] = pd.DataFrame([
        ['Target population', 'Human participants; population not specified in the reflection', 'Hypothetical second-year UB business-student prompt; no humans observed'],
        ['Pilot calibration', '16 participants in four chains; ordinary per-round times; 90-minute whole-game cap',
            f'{len(manifest.get("calibration_cases", []))} fixed states; {len(tables["calibration_checks"])} caps; {int(tables["calibration_checks"].n.sum())} selected state-cap observations'],
        ['Constraint', 'P25 restrictive / P95 generous / doubled generous seconds',
            f'Fixed high effort; generated-token caps {selection["short_tokens"]}/{selection["adequate_tokens"]}; calibration passed={selection["calibration_passed"]}'],
        ['Main sample', '80 new participants; ten chains per condition; pairing not specified', f'{selection["groups_per_condition"]} selected chain(s) per condition; repo pairs chains by block'],
        ['Information sharing', 'Shared supply-chain information', 'All four inventories; own demand, receipts, shipments, backlog, total supply line and history; other orders/backlogs hidden'],
        ['Process evidence', 'Screen recordings, written explanations and post-questionnaire', 'Exact prompts/replies/status/usage/state transitions; no human process evidence or questionnaire'],
    ], columns=['design_element', 'original_human_proposal', 'implemented_llm_study'])
    table('Original proposal and implemented study', 'design_mapping')
    text('Information sharing and mechanism coverage', [
        'The model sees all four current available inventories and its own local state/history, while other roles\' orders and backlogs remain hidden. '
        'Future customer demand is not disclosed. The same information fields and repeated question are supplied in both token conditions. '
        'Information load is not separately randomized in this pilot. The total outstanding supply line is already aggregated in the prompt, '
        'so the original proposal\'s mental aggregation across pipeline stages is not isolated. A token effect, if later observed, would not by itself identify that mechanism.'])
    table('H1 and H2 results', 'hypotheses', ['hypothesis', 'complete_pairs', 'estimate', 'predicted_direction',
          'p_raw', 'p_holm', 'ci_lower', 'ci_upper', 'status'])
    text('Individual reasons and hypothesis conclusions', [f'{r.hypothesis}: {r.individual_reasons.replace(".;", ";").rstrip(".")}. {r.conclusion}.'
        for r in tables['hypotheses'].itertuples(index=False)])
    table('Paired chain-block contrasts', 'block_contrasts', ['hypothesis', 'block', 'short_value',
          'adequate_value', 'contrast', 'usable_pair', 'pair_reason'])
    table('Inference prerequisites', 'inference_checks')
    text('Sample and data integrity', [
        f'The saved configuration requested at most {config["groups_per_condition"]} pairs, but selection.json prospectively froze '
        f'{selection["groups_per_condition"]} pair(s) using calibration charges before game outcomes. The selected count is the denominator, '
        'not the requested maximum. Each game has four roles, '
        f'{config["weeks"]} weeks, {config["warmup_weeks"]} fixed familiarization weeks and {config["continuation_weeks"]} automatic scoring weeks. '
        'Rules checks are common before treatment; passing them is a comprehension screen, not validation of human-like choice.',
        f'Saved run status: {progress["phase"]}; '
        f'valid model decisions: {saved_summary.get("valid_model_decisions")}; '
        f'missing decisions: {saved_summary.get("missing_decisions")}; fallbacks: {saved_summary.get("fallback_count")}. '
        'The group table records full-game and full-continuation availability. Invalid decisions or mechanical fallbacks are never labeled as model choices. '
        'All completed outcomes remain visible; no chain is filtered for cost, direction, zero orders or a favorable coefficient.',
        'The analysis independently rebuilds scored costs and checks saved outcomes, role fits, source provenance and raw requests. '
        'For this pilot, the separate independent_audit.json additionally reconstructs 320 physical game states, 288 model decisions and 96 continuation states. '
        'The week-12 filesystem recovery retained the same chains and saved decisions; it was not an extra replicate or behavioral replacement. '
        'Scheduling changed commonly from concurrency/interval 2/2 to 8/1; latency is consequently descriptive operational evidence, not a treatment or cognitive-time measurement. '
        'When exporting another run, rely on its inference checks and its own audit record rather than treating this pilot audit as transferable.'])
    table('Chain outcomes and H1 identifiability', 'groups', ['group_id', 'full_game', 'full_continuation',
        'valid_submitted_decisions', 'missing_decisions', 'fallback_count', 'total_scored_cost', 'usable_role_betas', 'group_beta'])
    if not completed_pilot:
        for kind, title, paragraphs in sections:
            if title == 'Sample and data integrity':
                paragraphs[2] = 'The analysis independently rebuilds scored costs and checks saved outcomes, role fits, source provenance and raw requests. '
                paragraphs[2] += 'Any independent audit and operational recovery records belong to this selected run. Incomplete games retain their observed costs; full-scored outcomes are missing until the scoring horizon finishes.'
    text('Token manipulation and technical timing', [
        f'Requested model: {config["model"]}; reasoning effort: {config["effort"]}, fixed across conditions. '
        f'Common technical timeout: {config["timeout_s"]} seconds. The cap covers total generated tokens, including hidden reasoning; '
        'it is an allowance, not computation necessarily consumed. The model is not told its cap. There is one structured decision call and no extra reasoning opportunity.',
        'The fixed calibration bank uses 24 matched hypothetical states per cap. Eligibility requires at least 95% valid decisions for both selected caps, '
        'recorded positive reasoning for valid answers, a complete matched bank, at least 25% lower median actual generated tokens at the short cap, '
        'and a generous reference with at most 5% of responses using at least 90% of its cap. Calibration usage summaries include incomplete generation. '
        'These are engineering feasibility thresholds, not validated human-time thresholds.',
        'In this pilot, 8,192 was the only candidate cap to reach 95% observed completion, but its median use was 5,057 versus 4,413 at the 16,384 reference '
        '(short use 14.59% higher). No pair passed the prespecified separation screen. The 8,192/16,384 diagnostic pair was declared before game outcomes; '
        'it was not chosen to favor H1/H2.',
        'In the completed games, mean output was 2,087.424 versus 2,183.375 tokens (4.39% lower in short), while medians were 1,982 versus 1,922 '
        '(3.12% higher in short). Both had zero calls reaching 90% of their assigned allowance. The short maximum was 7,258; adequate maximum 8,171. '
        'The game therefore supplies no clear binding-computation contrast. These per-call summaries are descriptive; 144 decisions per condition are not 144 independent subjects.',
        'A common transport timeout allows the API to finish a request and is not calibrated decision pressure. Adding a wall-clock treatment would require a separate '
        'prospective design accounting for server/network delay while holding token policy fixed. This dataset cannot estimate a human seconds-to-tokens conversion.'])
    if not completed_pilot:
        sections[-1][2][2:4] = ['Read the calibration and actual-token tables for this run. Diagnostic caps do not become a validated intervention merely because a game completes.']
    table('Calibration completion, truncation and actual computation', 'calibration_checks', ['limit', 'n',
        'completed_valid_n', 'completion_rate', 'median_output_tokens', 'near_cap_fraction', 'truncated_fraction', 'viable'])
    table('Main-game actual token usage and binding', 'token_use', ['condition', 'independent_chains',
        'decision_calls', 'measured_usage_calls', 'token_limit', 'median_output_tokens', 'mean_output_tokens',
        'median_reasoning_tokens', 'near_cap_fraction', 'cap_binding_fraction'])
    text('H1 structural coefficient checks', [
        'The inherited model is q_hat,t = max(0, F_t + alpha * (S - net_inventory_t - beta * supply_line_t)), '
        'F_t = theta * demand_t + (1-theta) * F_(t-1), initialized at 4. Each role is fitted on its own model-choice weeks. '
        'Bounds are theta/alpha in [0,1], beta in [0,2], target S in [0,10000], with multiple optimizer starts. '
        'Diagnostics check variation, sample size, inventory response, Jacobian rank/conditioning, poor fit and parameter limits. '
        'A boundary flag alone need not invalidate a fit; current upstream fits also have weak inventory response.',
        'For the 11 October pilot, retailer choices are constant at 8 in all model-choice weeks, so the required variation is absent. '
        'Other roles closely follow current incoming demand; fitted theta is approximately 1 and alpha is about 0.000003–0.000006, below the 0.02 response threshold. '
        'Their numerical beta values around 1.40–1.50 are unusable. Near-perfect R-squared measures reproduction of orders, not identification of supply-line weight.',
        'For an interior predicted order, the sensitivity to beta is -alpha * supply_line. As alpha approaches zero, changes in beta barely change predictions; '
        'at alpha=0 the model reduces to the forecast and beta drops out entirely. An unidentifiable beta is not beta=0 or proof of either neglect or full attention. '
        'All four roles must be usable before a chain mean enters H1; the current 0/8 usable fits leave zero eligible beta pairs. '
        'No selective averaging, new optimizer bounds or relaxed diagnostics were applied to produce a result.'])
    if not completed_pilot:
        sections[-1][2][1:] = ['The fit table records each role\'s diagnostics. Numerical coefficients flagged unusable are excluded from chain beta. '
            'Near-perfect R-squared does not establish identification: the sensitivity to beta is -alpha * supply_line for interior predictions, and beta drops out at alpha=0. '
            'No selective averaging or relaxed diagnostics are applied.']
    table('All role fits, including unusable numerical coefficients', 'beta_fits', ['condition', 'role',
        'theta', 'alpha', 'beta', 'rmse', 'r2', 'jac_condition', 'usable', 'fit_flag'])
    text('H2 costs and descriptive effect', [
        'Costs are rebuilt from physical inventory and backlog, including the full fixed scoring horizon. '
        'The table separates common familiarization, model-choice weeks and automatic continuation. '
        'The completed pilot has 96 familiarization cost plus 60 model-choice cost and zero continuation cost per chain; all cost is holding, with zero backlog. '
        'The paired condition difference is 156 - 156 = 0, or 0% relative to adequate. These are simulated team costs, separate from actual API charges. '
        'H1 coefficient failure does not remove this cost observation.',
        'There is only one independent pair. The between-pair variance, standard error and population confidence interval are not estimable here. '
        'The reported zero is an observed difference, not a statistically established zero effect, a noninferiority result or evidence that H2 is false. '
        'The failed computation manipulation is an additional reason that this cannot adjudicate a binding-constraint hypothesis.'])
    if not completed_pilot:
        sections[-1][2][:] = ['Costs are rebuilt from inventory and backlog through the full scoring horizon. H1 coefficient failure does not remove a complete H2 cost observation. '
            'Use the paired table for the observed difference and eligibility checks for inference. No one-pair confidence interval or equivalence claim is produced.']
    table('Holding and backlog costs by scoring phase', 'cost_components')
    text('Bullwhip and order-behavior checks: supplementary outcomes', [
        'Direct bullwhip ratios supplement H1/H2; they are not a newly declared third confirmatory hypothesis. '
        'The primary descriptive window includes all game weeks, as in the existing repo workflow. '
        'Local ratio = sample Var(role orders) / sample Var(local incoming demand); customer ratio uses the common retailer demand denominator. '
        'The window and numerator/denominator variances are explicit. A constant denominator is missing, never zero or infinity.',
        'In this pilot all local ratios are 1.000. Customer ratios are 1.000, 1.215278, 1.416667 and 1.604167 upstream. '
        'All matched orders agree, with values only 4 or 8, and each equals its role\'s incoming demand. '
        'The >1 customer ratios arise because the same demand step is delayed upstream within a finite window. '
        'They do not show local amplification or sustained oscillations. They also do not imply that the student persona is realistic.',
        'The post-hoc window sensitivity table adds model-choice weeks (5 onward) and a settled-period window (9 onward). '
        'After familiarization, customer demand is constant so customer-referenced ratios are undefined. In the settled window, local denominators are also constant. '
        'These windows illustrate the measurement limitation; they were not used to select caps, exclude chains or seek a favorable hypothesis result. '
        'Zero orders remain valid outcomes and no advertised maximum or positive-order rule was imposed.'])
    if not completed_pilot:
        sections[-1][2][1] = 'The agreement table compares matched valid decisions across conditions. Full-window and sensitivity ratios are descriptive; denominator variance and missing values remain visible.'
        sections[-1][2][2] = 'The post-hoc sensitivity table includes model-choice weeks and a late window starting at week 9. Those windows can be undefined if the denominator is constant; no settling assumption is imposed on this selected run. They do not determine cap selection, chain exclusions or hypothesis eligibility.'
    table('Matched decisions across conditions', 'order_agreement')
    table('Role ordering, backlog and service checks', 'behavior_checks', ['condition', 'role', 'valid_choices',
        'zero_orders', 'max_order', 'concerning_zeros', 'backlog_weeks', 'max_backlog', 'cumulative_shipped_over_demand'])
    table('Role holding costs and terminal-state checks', 'behavior_checks', ['condition', 'role',
        'observed_holding_cost', 'observed_backlog_cost', 'terminal_inventory', 'terminal_backlog', 'terminal_preorder_supply_line'])
    text('Service and terminal-state definitions', [
        'Cumulative shipped divided by demand can include late delivery; it is not inherently an on-time service rate. '
        'Backlog weeks and maximum backlog are therefore shown separately. In this pilot there was no backlog. '
        'Final inventory and backlog are zero, but the saved final pre-order supply line is 16 cases per role; outstanding orders have not disappeared. '
        'Continuation is a fixed finite scoring policy, not an empty-system or infinite-horizon terminal valuation.'])
    if not completed_pilot:
        sections[-1][2][:] = ['Cumulative shipped divided by demand can include late delivery and is not inherently an on-time service rate. Backlog weeks and maximum backlog are shown separately. The terminal pre-order supply line remains visible; finite continuation does not imply an empty system or infinite-horizon terminal valuation.']
    table('Bullwhip window sensitivity with denominator variances', 'bullwhip_windows', ['condition', 'role',
        'first_week', 'last_week', 'order_variance', 'local_demand_variance', 'customer_demand_variance',
        'bullwhip_vs_local', 'bullwhip_vs_customer'])
    sections.append(('image', 'Order and local-demand trajectories', 'order_demand_checks.png'))
    sections.append(('image', 'Actual token distributions', 'actual_token_ecdf.png'))
    text('Interpretation, power and the next study', [
        'H1 is not testable with these unusable structural coefficients. H2 has an observed zero difference with insufficient independent replication. '
        'The computational treatment was not validated. More copies of the same easy step trajectories would not repair an unidentifiable coefficient or a nonbinding intervention.',
        'Before an inferential study, prospectively specify an operationally meaningful effect, a demand/state design that permits coefficient identification, '
        'and a calibration that can separate actual computation while preserving decision completion. '
        'Use offline synthetic recovery checks for software and identifiability, then a separately authorized live feasibility check for the intervention. '
        'Any changed demand, prompt, model or cap policy creates a new scientific protocol/fingerprint. Do not retune it for the sign or significance of H1/H2.',
        'The original reflection proposed ten chains per condition. Neither that target nor the inherited eight-pair floor guarantees power. '
        'This one-pair pilot cannot estimate between-pair variance or support a defensible numeric power claim. '
        'A future sample plan needs independent-pair variance and a meaningful effect, not the variance of repeated weekly calls. '
        'Any budget/power tradeoff should be declared before outcomes; $10 funded this diagnostic pilot, not a powered hypothesis test.',
        'A hypothetical UB student prompt is a simulation condition, not population validation. Compare at least a small set of matched real-student decisions '
        'before making claims about actual students, using a predefined comparison of choices and trajectories. '
        'This run did not test the 18 behavioral biases mentioned in the research background, collect human observations or isolate an information-overload mechanism. '
        'No bias was inserted to make the model appear human.',
        'Historical none-effort/two-call nano and GPT-5.6 Luna results remain interpretable in their own reports. They cannot be pooled with this high-effort/single-call study '
        'or used as a causal estimate of the effect of the new persona, model or effort, because several features changed together.'])
    if not completed_pilot:
        sections[-1][2][0] = 'Read the hypothesis and eligibility tables for this selected run. Inadequate replication, failed manipulation, or coefficient nonidentifiability each limit interpretation. Repeating an unidentifiable or nonbinding setup does not repair its design.'
        sections[-1][2][2] = 'The original reflection proposed ten chains per condition. Neither that target nor the inherited eight-pair floor guarantees power. A future sample plan needs independent-pair variability and a meaningful effect, not variance across repeated weekly calls. Budget and sample selection must be declared before outcomes.'
    text('Exact prompt, model and reproducibility', [
        f'Scientific protocol: {manifest["protocol"]}. Run: {run.name}. Analysis version: {VERSION}. '
        f'Analysis fingerprint: {fingerprint}. The unchanged scientific source hashes are in the run manifest; '
        'new analysis-source snapshots and hashes are recorded separately in analysis_manifest.json.',
        f'The total saved campaign ledger is ${saved_summary.get("estimated_charged_or_reserved_usd", saved_summary.get("estimated_charge_usd", 5.19360675)):.8f} '
        '(see the original report summary and ledger for accounting scope). This analysis makes zero API calls and adds zero API charges. '
        'For the 11 October campaign, calibration cost $2.24484225, including preserved excluded old-state charges; rules cost $0.013455 and game decisions $2.93530950. '
        'Copied/imported evidence is charged once and is not an independent observation.',
        'Exact rendered messages, histories, schema, requested/returned model, high effort, cap, completion status, usage, latency and charge remain in requests/*.json. '
        'prompt_template.txt below reproduces the manifest template, while raw requests are authoritative for each realized prompt. '
        'The report directory contains hypothesis, contrast, eligibility, token, cost, fit and bullwhip CSVs; the workbook contains the same tables. '
        'analysis_manifest.json records byte hashes of every input, including all raw request files and this report\'s analysis sources.',
        'Rebuild from PowerShell in D:\\beergame: .\\.venv\\Scripts\\python.exe -X utf8 scripts/build_student_hypothesis_report.py --run ' + str(run),
        'Notebook 03 includes a separate full-hypothesis review cell after the game/report cells. It reads saved evidence and needs no key. '
        'Completed raw evidence and the five pinned scientific source files are preserved; this report has a separate analysis fingerprint.'])
    if not completed_pilot:
        sections[-1][2][1] = 'This analysis makes zero API calls and adds zero API charges. Original run reports and request ledgers record spending and its accounting scope; copied evidence is not another observation.'
    for filename, frame in tables.items():
        frame.to_csv(output / f'{filename}.csv', index=False)
    plots(run, output)
    lines = ['# Full hypothesis analysis: student Beer Game', '', f'Analysis date: {analysis_date}. LLM simulations; no human observations.', '']
    for kind, title, payload in sections:
        lines += ['## ' + title, '']
        if kind == 'text':
            for paragraph in payload:
                lines += [paragraph, '']
        elif kind == 'table':
            lines += ['| ' + ' | '.join(payload.columns) + ' |', '| ' + ' | '.join(['---'] * len(payload.columns)) + ' |']
            lines += ['| ' + ' | '.join(show(v) for v in row) + ' |' for row in payload.itertuples(index=False, name=None)]
            lines += ['']
        else:
            lines += [f'![{title}]({payload})', '']
    (output / MARKDOWN_REPORT).write_text('\n'.join(lines), encoding='utf-8')
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document(); doc.add_heading('Full hypothesis analysis: student Beer Game', 0)
    doc.add_paragraph(f'{analysis_date} • LLM simulations; no human observations')
    doc.styles['Normal'].font.size = Pt(9)
    section = doc.sections[0]; section.page_width = Inches(11.7); section.page_height = Inches(8.3)
    section.left_margin = section.right_margin = Inches(.5)
    for kind, title, payload in sections:
        doc.add_heading(title, 1)
        if kind == 'text':
            for paragraph in payload: doc.add_paragraph(paragraph)
        elif kind == 'table':
            t = doc.add_table(rows=1, cols=len(payload.columns)); t.style = 'Light Shading Accent 1'
            for cell, col in zip(t.rows[0].cells, payload.columns): cell.text = col.replace('_', ' ')
            for row in payload.itertuples(index=False, name=None):
                for cell, value in zip(t.add_row().cells, row): cell.text = show(value)
            for row in t.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for item in paragraph.runs: item.font.size = Pt(6 if len(payload.columns) > 10 else 8)
        else:
            doc.add_picture(str(output / payload), width=Inches(8))
    doc.save(output / WORD_REPORT)
    with pd.ExcelWriter(output / WORKBOOK, engine='openpyxl') as writer:
        for name, frame in tables.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
        pd.DataFrame([dict(input_file=k, sha256=v) for k, v in input_hashes.items()]).to_excel(writer, sheet_name='Input_hashes', index=False)
        pd.DataFrame([dict(source=k, sha256=v) for k, v in sources.items()]).to_excel(writer, sheet_name='Analysis_sources', index=False)
    prompt = '\n\n'.join(f'{row["role"]}:\n{row["content"]}' for row in manifest['question_template'])
    (output / 'prompt_template.txt').write_text(prompt + '\n', encoding='utf-8')
    final_inputs = {name: sha(run / name) for name in inputs}
    final_calibration = {str(p.relative_to(calibration_source)): sha(p) for p in calibration_inputs}
    if final_inputs != input_hashes or final_calibration != identity['calibration_evidence']['file_sha256']:
        raise RuntimeError('Saved evidence changed during analysis; report is not certified.')
    provenance = dict(identity, analysis_fingerprint=fingerprint, input_evidence_unchanged=True,
                      run_directory=str(run), report_directory=str(output),
                      generated_utc=datetime.now(timezone.utc).isoformat(),
                      paid_api_calls=0, additional_api_charge_usd=0,
                      analysis_summary=plain(result.get('summary', {})),
                      analysis_checks=plain(result.get('provenance', {})))
    (output / 'analysis_manifest.json').write_text(json.dumps(plain(provenance), indent=2, allow_nan=False), encoding='utf-8')
    pointer = dict(run_directory=str(run), analysis_directory=str(output), fingerprint=fingerprint,
                   markdown_report=str(output / MARKDOWN_REPORT), word_report=str(output / WORD_REPORT),
                   workbook=str(output / WORKBOOK), paid_api_calls=0)
    if update_pointer:
        (reports_root / 'LATEST_STUDENT_HYPOTHESIS_REPORT.json').write_text(json.dumps(pointer, indent=2), encoding='utf-8')
        (reports_root / 'LATEST_REPORT.txt').write_text('\n'.join(pointer[k] for k in
            ('word_report', 'markdown_report', 'workbook')) + '\n', encoding='utf-8')
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run', required=True, type=Path)
    args = parser.parse_args()
    print(build_report(args.run))


if __name__ == '__main__':
    main()
