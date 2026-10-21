"""Evidence-based reports for the isolated student-persona protocol.

Importing this module never creates an API client or modifies historical data.
The report describes observed simulations and retains missing decisions/fallbacks.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd

from .game import ROLES


SOURCES = [
    ("MIT / John Sterman: Beer Game", "https://web.mit.edu/jsterman/www/SDG/beergame.html"),
    ("OpenAI: reasoning models and output-token accounting", "https://developers.openai.com/api/docs/guides/reasoning"),
    ("OpenAI: structured outputs", "https://developers.openai.com/api/docs/guides/structured-outputs"),
    ("Bisbee et al.: limitations of LLM survey respondents", "https://doi.org/10.1017/pan.2024.5"),
    ("GPT and behavioral operations study", "https://doi.org/10.1287/msom.2023.0279"),
    ("OpenAI: GPT-5.4 Mini model documentation", "https://developers.openai.com/api/docs/models/gpt-5.4-mini"),
]


def _json(path, default):
    path = Path(path)
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else default


def _csv(path):
    path = Path(path)
    if not path.is_file():
        return pd.DataFrame()
    try:
        return pd.read_csv(path)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def _finite(value):
    """Make summaries valid JSON, including estimator NaN/infinite diagnostics."""
    if isinstance(value, dict):
        return {str(k): _finite(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [_finite(v) for v in value]
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value) if math.isfinite(float(value)) else None
    if value is pd.NA:
        return None
    return value


def _bool(frame, name, default=False):
    if name not in frame:
        return pd.Series(default, index=frame.index, dtype=bool)
    return frame[name].map(lambda x: x is True or str(x).lower() in ("true", "1", "1.0"))


def _number(frame, name, default=0):
    if name not in frame:
        return pd.Series(default, index=frame.index, dtype=float)
    return pd.to_numeric(frame[name], errors="coerce").fillna(default)


def _measured(frame, synthetic=False):
    if synthetic:
        return frame
    if "usage_observed" in frame:
        return frame[_bool(frame, "usage_observed")]
    if "output_tokens" in frame and "reasoning_tokens" in frame:
        return frame[frame.output_tokens.notna() & frame.reasoning_tokens.notna()]
    return frame.iloc[:0]


def _variance(series):
    values = np.asarray(series, dtype=float)
    return float(np.var(values, ddof=1)) if len(values) > 1 else np.nan


def _ratio(numerator, denominator):
    den = _variance(denominator)
    return _variance(numerator) / den if np.isfinite(den) and den > 1e-12 else np.nan


def _format(value):
    if value is None or value is pd.NA:
        return "missing"
    if isinstance(value, (float, np.floating)):
        return f"{value:,.4g}" if np.isfinite(value) else "missing"
    if isinstance(value, (bool, np.bool_)):
        return "yes" if value else "no"
    return str(value).replace("|", "\\|").replace("\n", " ")


def _table(frame, columns=None, limit=40):
    if frame.empty:
        return "No observations recorded."
    if columns is not None:
        frame = frame[[c for c in columns if c in frame]]
    shown = frame.head(limit)
    lines = ["| " + " | ".join(map(str, shown.columns)) + " |",
             "| " + " | ".join(["---"] * len(shown.columns)) + " |"]
    lines.extend("| " + " | ".join(_format(v) for v in row) + " |"
                 for row in shown.itertuples(index=False, name=None))
    if len(frame) > limit:
        lines += ["", f"Showing {limit} of {len(frame)} rows; the workbook and CSV contain all rows."]
    return "\n".join(lines)


def _request_table(directory, config):
    rows = []
    for path in sorted((directory / "requests").glob("*.json")):
        record = _json(path, {})
        identity = record.get("identity") or {}
        raw = record.get("response") or {}
        usage = record.get("usage") or raw.get("usage") or {}
        details = usage.get("output_tokens_details") or {}
        synthetic = bool(record.get("synthetic", config.get("mode") == "offline"))
        usage_observed = bool(record.get("usage_observed", bool(usage) and not synthetic))
        measured = usage_observed or synthetic
        reasoning = usage.get("reasoning_tokens", details.get("reasoning_tokens")) if measured else None
        output = usage.get("output_tokens") if measured else None
        try:
            answer = json.loads(record.get("text", ""))
        except (TypeError, ValueError):
            answer = None
        valid_order = (record.get("status") == "ok" and record.get("response_status") == "completed"
                       and not record.get("refused") and isinstance(answer, dict)
                       and set(answer) == {"order"} and type(answer["order"]) is int
                       and 0 <= answer["order"] <= config.get("emergency_order_limit", 100000))
        rows.append(dict(
            request_file=path.name, group_id=identity.get("group", identity.get("group_id", "")),
            role=identity.get("role", ""), week=identity.get("week"),
            stage=identity.get("stage", ""), condition=identity.get("condition", ""),
            requested_model=record.get("requested_model", record.get("model", config.get("model", "unknown"))),
            returned_model="" if synthetic else record.get("returned_model", raw.get("model", "")),
            effort=config.get("effort", config.get("reasoning_effort", "unknown")),
            status=record.get("status", "unknown"), response_status=record.get("response_status", raw.get("status", "unknown")),
            incomplete_reason=record.get("incomplete_reason", (raw.get("incomplete_details") or {}).get("reason", "")),
            valid_order_submission=valid_order,
            submitted_order=answer.get("order") if isinstance(answer, dict) and type(answer.get("order")) is int else None,
            usage_observed=usage_observed,
            input_tokens=usage.get("input_tokens") if measured else None,
            output_tokens=output, reasoning_tokens=reasoning,
            nonreasoning_tokens=output - reasoning if output is not None and reasoning is not None else None,
            token_limit=record.get("limit", record.get("token_limit", record.get("max_output_tokens"))),
            elapsed_s=record.get("elapsed_s", 0), charge_usd=record.get("charge_usd", 0),
            synthetic=synthetic,
        ))
    return pd.DataFrame(rows)


def _game_tables(rounds, tail, config):
    groups, roles, fits = [], [], []
    expected_weeks = int(config.get("weeks", 40))
    if rounds.empty:
        return pd.DataFrame(), pd.DataFrame(), pd.DataFrame()
    for gid, game in rounds.groupby("group_id"):
        game = game.sort_values(["week", "role"])
        meta = game.iloc[0]
        observed_weeks = int(game.week.nunique())
        complete = observed_weeks == expected_weeks and len(game) == expected_weeks * len(ROLES)
        agent_game = game[game.phase == "agent"]
        fallback_count = int(_bool(agent_game, "fallback").sum())
        invalid_count = int((~_bool(agent_game, "valid_decision")).sum())
        t = tail[tail.group_id == gid] if not tail.empty and "group_id" in tail else pd.DataFrame()
        observed_cost = float(_number(game, "cost").sum())
        continuation_cost = float(_number(t, "cost").sum()) if not t.empty else np.nan
        tail_complete = not t.empty and t.week.nunique() == int(config.get("continuation_weeks", 12)) and len(t) == int(config.get("continuation_weeks", 12)) * len(ROLES)
        retailer = game[game.role == "retailer"].set_index("week")
        for role, trajectory in game.groupby("role"):
            trajectory = trajectory.sort_values("week")
            agent = trajectory[trajectory.phase == "agent"]
            valid = agent[_bool(agent, "valid_decision") & ~_bool(agent, "fallback")]
            measured = _measured(agent, config.get("mode") == "offline")
            customer = retailer.reindex(trajectory.week).demand.to_numpy() if "demand" in retailer else np.array([])
            local_variance = _variance(trajectory.demand)
            customer_variance = _variance(customer)
            roles.append(dict(
                group_id=gid, block=int(meta.get("block", 0)), condition=meta.get("condition", "unknown"), role=role,
                observed_weeks=len(trajectory), full_game=complete,
                committed_agent_orders=len(agent), valid_submitted_decisions=len(valid),
                fallback_count=int(_bool(agent, "fallback").sum()),
                missing_decisions=int((~_bool(agent, "valid_decision")).sum()),
                zero_orders_valid=int((valid.order == 0).sum()) if not valid.empty else 0,
                zero_rate_valid=float((valid.order == 0).mean()) if not valid.empty else np.nan,
                zero_rate_committed=float((agent.order == 0).mean()) if not agent.empty else np.nan,
                max_submitted_order=float(valid.order.max()) if not valid.empty else np.nan,
                max_committed_order=float(agent.order.max()) if not agent.empty else np.nan,
                concerning_zero_orders=int(((valid.order == 0) & (valid.inventory == 0) & (valid.backlog > 0) & (valid.supply_line == 0)).sum()) if not valid.empty else 0,
                order_variance=_variance(trajectory.order), local_demand_variance=local_variance,
                customer_demand_variance=customer_variance,
                bullwhip_vs_local=_ratio(trajectory.order, trajectory.demand),
                bullwhip_vs_customer=_ratio(trajectory.order, customer),
                observed_holding_cost=float((.5 * _number(trajectory, "inventory")).sum()),
                observed_backlog_cost=float(_number(trajectory, "backlog").sum()),
                measured_usage_decisions=len(measured),
                output_tokens=int(_number(measured, "output_tokens").sum()),
                reasoning_tokens=int(_number(measured, "reasoning_tokens").sum()),
                nonreasoning_tokens=int(_number(measured, "nonreasoning_tokens").sum()),
            ))
            fit = dict(theta=np.nan, alpha=np.nan, beta=np.nan, target=np.nan, rmse=np.nan,
                       r2=np.nan, jac_condition=np.nan, usable=False, n=len(agent), fit_flag="missing_decisions_or_fallback")
            if not agent.empty and not _bool(agent, "fallback").any() and _bool(agent, "valid_decision").all():
                try:
                    from .estimation import estimate_beta
                    fit = estimate_beta(agent)
                    if not complete:
                        fit["usable"] = False
                        fit["fit_flag"] += ";partial_game"
                    if int((agent.order > 0).sum()) < 5:
                        fit["usable"] = False
                        fit["fit_flag"] += ";too_few_positive_orders"
                except ImportError:
                    fit["fit_flag"] = "estimator_dependency_unavailable"
            fits.append(dict(group_id=gid, block=int(meta.get("block", 0)), condition=meta.get("condition", "unknown"), role=role, **fit))
        group_fits = [fit for fit in fits if fit["group_id"] == gid]
        final = t[t.week == t.week.max()] if not t.empty else game[game.week == game.week.max()]
        groups.append(dict(
            group_id=gid, block=int(meta.get("block", 0)), condition=meta.get("condition", "unknown"),
            observed_weeks=observed_weeks, full_game=complete, full_continuation=tail_complete,
            valid_submitted_decisions=int((_bool(agent_game, "valid_decision") & ~_bool(agent_game, "fallback")).sum()),
            missing_decisions=invalid_count, fallback_count=fallback_count,
            observed_game_cost=observed_cost, continuation_cost=continuation_cost,
            total_observed_cost=observed_cost + (continuation_cost if np.isfinite(continuation_cost) else 0),
            total_scored_cost=observed_cost + continuation_cost if complete and tail_complete else np.nan,
            clean_model_trajectory=complete and tail_complete and invalid_count == 0 and fallback_count == 0,
            usable_role_betas=sum(bool(fit["usable"]) for fit in group_fits),
            group_beta=float(np.mean([fit["beta"] for fit in group_fits])) if len(group_fits) == len(ROLES) and all(fit["usable"] for fit in group_fits) else np.nan,
            terminal_inventory=float(_number(final, "inventory").sum()),
            terminal_backlog=float(_number(final, "backlog").sum()),
        ))
    return pd.DataFrame(groups), pd.DataFrame(roles), pd.DataFrame(fits)


def _condition_table(groups, rounds, synthetic=False):
    if groups.empty:
        return pd.DataFrame()
    rows = []
    for condition, own in groups.groupby("condition"):
        agent = rounds[(rounds.condition == condition) & (rounds.phase == "agent")]
        valid = agent[_bool(agent, "valid_decision") & ~_bool(agent, "fallback")]
        measured = _measured(agent, synthetic)
        rows.append(dict(
            condition=condition, chains_observed=len(own), chains_complete=int((own.full_game & own.full_continuation).sum()),
            clean_chains=int(own.clean_model_trajectory.sum()), valid_submitted_decisions=len(valid),
            missing_decisions=int(own.missing_decisions.sum()), fallback_count=int(own.fallback_count.sum()),
            zero_rate_valid=float((valid.order == 0).mean()) if not valid.empty else np.nan,
            max_submitted_order=float(valid.order.max()) if not valid.empty else np.nan,
            mean_total_scored_cost=float(own.total_scored_cost.mean()) if own.total_scored_cost.notna().any() else np.nan,
            measured_usage_decisions=len(measured),
            mean_output_tokens=float(_number(measured, "output_tokens").mean()) if not measured.empty else np.nan,
            mean_reasoning_tokens=float(_number(measured, "reasoning_tokens").mean()) if not measured.empty else np.nan,
        ))
    return pd.DataFrame(rows)


def _plots(directory, rounds, tail, requests, synthetic):
    if rounds.empty and requests.empty:
        return []
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    label = "SYNTHETIC SOFTWARE CHECK — " if synthetic else "LLM feasibility pilot — "
    files = []
    colors = {"short": "#b83c39", "adequate": "#247799"}
    if not rounds.empty:
        figure, axes = plt.subplots(2, 2, figsize=(12, 7), sharex=True)
        for axis, role in zip(axes.flat, ROLES):
            for gid, own in rounds[rounds.role == role].groupby("group_id"):
                own = own.sort_values("week")
                color = colors.get(own.condition.iloc[0], "#777777")
                axis.plot(own.week, own.order, label=gid, color=color, alpha=.7)
                fallback = own[_bool(own, "fallback")]
                if not fallback.empty:
                    axis.scatter(fallback.week, fallback.order, marker="x", color="black", zorder=5)
            axis.set_title(role.title()); axis.set_xlabel("Week"); axis.set_ylabel("Cases ordered")
            axis.grid(alpha=.2); axis.legend(fontsize=7)
        figure.suptitle(label + "committed orders (black × = fallback)")
        figure.tight_layout(); figure.savefig(directory / "orders_by_role.png", dpi=160); plt.close(figure)
        files.append("orders_by_role.png")
        figure, axis = plt.subplots(figsize=(10, 5))
        for gid, own in rounds.groupby("group_id"):
            t = tail[tail.group_id == gid] if not tail.empty else pd.DataFrame()
            values = pd.concat([own[["week", "cost"]], t[["week", "cost"]] if not t.empty else own.iloc[:0][["week", "cost"]]])
            costs = values.groupby("week").cost.sum().sort_index().cumsum()
            axis.plot(costs.index, costs.values, label=gid, color=colors.get(own.condition.iloc[0]), alpha=.75)
        axis.set_title(label + "cumulative observed team cost")
        axis.set_xlabel("Week, including automatic continuation"); axis.set_ylabel("Holding + backlog cost")
        axis.grid(alpha=.2); axis.legend(fontsize=8); figure.tight_layout()
        figure.savefig(directory / "cumulative_cost.png", dpi=160); plt.close(figure)
        files.append("cumulative_cost.png")
    if not requests.empty:
        figure, axis = plt.subplots(figsize=(10, 4))
        plotted = requests[requests.stage.isin(["decision", "order", "calibration"])]
        if plotted.empty:
            plotted = requests
        axis.scatter(np.arange(len(plotted)), pd.to_numeric(plotted.output_tokens, errors="coerce"), s=14, alpha=.65, label="Total output")
        axis.scatter(np.arange(len(plotted)), pd.to_numeric(plotted.reasoning_tokens, errors="coerce"), s=12, alpha=.65, label="Reasoning subset")
        axis.set_title(label + "recorded token usage")
        axis.set_xlabel("Saved request index (file order)"); axis.set_ylabel("Tokens")
        axis.grid(alpha=.2); axis.legend(); figure.tight_layout()
        figure.savefig(directory / "token_usage.png", dpi=160); plt.close(figure)
        files.append("token_usage.png")
    return files


def _workbook(directory, tables, summary):
    import openpyxl  # noqa: F401 — optional export dependency
    with pd.ExcelWriter(directory / "DETAILS.xlsx", engine="openpyxl") as writer:
        pd.DataFrame([dict(metric=k, value=json.dumps(_finite(v)) if isinstance(v, (list, dict)) else _finite(v))
                      for k, v in summary.items()]).to_excel(writer, sheet_name="Summary", index=False)
        for name, frame in tables.items():
            export = frame.copy()
            for column in export:
                export[column] = export[column].map(lambda v: json.dumps(_finite(v)) if isinstance(v, (dict, list)) else v)
            export.to_excel(writer, sheet_name=name[:31], index=False)
        for sheet in writer.book:
            sheet.freeze_panes = "A2"
            if sheet.max_row > 1:
                sheet.auto_filter.ref = sheet.dimensions
            for cells in sheet.columns:
                sheet.column_dimensions[cells[0].column_letter].width = min(48, max(12, len(str(cells[0].value)) + 2))


def _docx(directory, paragraphs, tables, plots):
    from docx import Document
    from docx.shared import Inches
    document = Document()
    document.add_heading("Student-persona Beer Game: evidence report", 0)
    for title, text in paragraphs:
        if title:
            document.add_heading(title, 1)
        for paragraph in text.split("\n\n"):
            document.add_paragraph(paragraph)
    for name in ("Condition outcomes", "Group outcomes", "Calibration", "Beta diagnostics"):
        frame = tables[name]
        if frame.empty:
            continue
        document.add_heading(name, 1)
        wanted = {
            "Condition outcomes": ["condition", "chains_complete", "clean_chains", "missing_decisions", "fallback_count", "zero_rate_valid", "max_submitted_order", "mean_total_scored_cost"],
            "Group outcomes": ["group_id", "full_game", "fallback_count", "total_scored_cost", "usable_role_betas"],
            "Calibration": list(frame.columns)[:8],
            "Beta diagnostics": ["group_id", "role", "beta", "usable", "fit_flag"],
        }[name]
        shown = frame[[c for c in wanted if c in frame]].head(40)
        table = document.add_table(rows=1, cols=len(shown.columns)); table.style = "Light Shading Accent 1"
        for cell, name in zip(table.rows[0].cells, shown.columns):
            cell.text = name
        for row in shown.itertuples(index=False, name=None):
            for cell, value in zip(table.add_row().cells, row):
                cell.text = _format(value)
    for name in plots:
        document.add_picture(str(directory / name), width=Inches(6.3))
    document.save(directory / "REPORT.docx")


def build_report(directory):
    """Write a reviewable report, including stopped/partial and offline runs.

    Returns the run directory. Optional Word/Excel/plot dependencies may be absent;
    Markdown and a strict-JSON summary are still produced with explicit warnings.
    """
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    manifest = _json(directory / "manifest.json", {})
    config = manifest.get("config") or {}
    progress = _json(directory / "progress.json", {})
    selection = _json(directory / "selection.json", {})
    recommendation = _json(directory / "calibration_recommendation.json", {})
    calibration = _csv(directory / "calibration.csv")
    synthetic = config.get("mode") == "offline"
    calibration_summary = pd.DataFrame(recommendation.get("screen_rows", []))
    if calibration_summary.empty and not calibration.empty and "cap" in calibration:
        summarized = []
        for cap, own in calibration.groupby("cap"):
            valid = _bool(own, "valid_decision")
            measured = _measured(own, synthetic)
            summarized.append(dict(cap=cap, n=len(own), valid_decisions=int(valid.sum()),
                                   completion_rate=float(valid.mean()),
                                   measured_usage_n=len(measured),
                                   mean_output_tokens=float(_number(measured, "output_tokens").mean()) if not measured.empty else np.nan,
                                   mean_reasoning_tokens=float(_number(measured, "reasoning_tokens").mean()) if not measured.empty else np.nan))
        calibration_summary = pd.DataFrame(summarized)
    rounds, tail = _csv(directory / "rounds.csv"), _csv(directory / "continuation.csv")
    rules = _json(directory / "rules_checks.json", [])
    rules_frame = pd.json_normalize(rules if isinstance(rules, list) else rules.get("checks", []))
    requests = _request_table(directory, config)
    groups, roles, fits = _game_tables(rounds, tail, config)
    conditions = _condition_table(groups, rounds, synthetic)
    expected_groups = 2 * int(config.get("groups_per_condition", config.get("groups", 3)))
    phase = progress.get("phase", progress.get("status", "unknown"))
    calibration_passed = bool(recommendation.get("manipulation_screen_passed", recommendation.get("passed", False)))
    complete_count = int((groups.full_game & groups.full_continuation).sum()) if not groups.empty else 0
    clean_count = int(groups.clean_model_trajectory.sum()) if not groups.empty else 0
    fallback_count = int(groups.fallback_count.sum()) if not groups.empty else 0
    missing_count = int(groups.missing_decisions.sum()) if not groups.empty else 0
    returned_models = sorted(set(requests.returned_model.dropna()) - {""}) if not requests.empty else []
    decision_requests = requests[requests.stage == "decision"] if not requests.empty else pd.DataFrame()
    clean_study = not synthetic and calibration_passed and complete_count == expected_groups and clean_count == expected_groups and phase == "complete"
    summary = dict(
        protocol=manifest.get("protocol", "unknown"), mode=config.get("mode", "unknown"), synthetic=synthetic,
        status=phase, requested_model=config.get("model", "unknown"), returned_models=returned_models,
        effort=config.get("effort", config.get("reasoning_effort", "unknown")),
        calibration_passed=calibration_passed, expected_chains=expected_groups, completed_chains=complete_count,
        clean_chains=clean_count, valid_model_decisions=int(groups.valid_submitted_decisions.sum()) if not groups.empty else 0,
        valid_submitted_decisions=int(groups.valid_submitted_decisions.sum()) if not groups.empty else 0,
        valid_synthetic_decisions=int(groups.valid_submitted_decisions.sum()) if synthetic and not groups.empty else 0,
        missing_decisions=missing_count, fallback_count=fallback_count, clean_behavior_study=clean_study,
        saved_failed_decision_requests=int((~_bool(decision_requests, "valid_order_submission")).sum()),
        unresolved_or_unmeasured_usage_requests=int((~_bool(requests, "usage_observed") & ~_bool(requests, "synthetic")).sum()),
        saved_requests=len(requests), estimated_charged_or_reserved_usd=float(_number(requests, "charge_usd").sum()),
        input_tokens=int(_number(requests, "input_tokens").sum()), output_tokens=int(_number(requests, "output_tokens").sum()),
        reasoning_tokens=int(_number(requests, "reasoning_tokens").sum()),
        nonreasoning_tokens=int(_number(requests, "nonreasoning_tokens").sum()),
        selected_caps=selection or manifest.get("chosen_caps", config.get("chosen_caps", {})),
        usable_role_betas=int(_bool(fits, "usable").sum()), total_role_beta_fits=len(fits),
        confirmatory_p_values=False,
    )
    if synthetic:
        summary["valid_model_decisions"] = 0
    evidence = ("SYNTHETIC SOFTWARE CHECK: no actual model responses, human observations or treatment-effect evidence."
                if synthetic else "Saved API evidence from hypothetical LLM role agents; no real student responses were collected.")
    outcome = ("All planned games and continuation weeks completed without substituted decisions, and the calibration screen passed. Results remain exploratory."
               if clean_study else f"A clean behavioral comparison is not established: calibration passed={calibration_passed}, completed chains={complete_count}/{expected_groups}, clean chains={clean_count}, missing decisions={missing_count}, fallback orders={fallback_count}.")
    stop = progress.get("reason", progress.get("stop_reason", ""))
    game_weeks = int(config.get("weeks", 40))
    warmup_weeks = int(config.get("warmup_weeks", 4))
    continuation_weeks = int(config.get("continuation_weeks", 12))
    timeout = config.get("timeout_s", 180)
    selected_short = selection.get("short_tokens", recommendation.get("short_tokens", config.get("short_tokens")))
    selected_adequate = selection.get("adequate_tokens", recommendation.get("adequate_tokens", config.get("adequate_tokens")))
    if calibration_passed:
        calibration_statement = f"Calibration selected total-output caps {selected_short}/{selected_adequate} (short/adequate); these are provisional pilot thresholds."
    elif selection:
        calibration_statement = f"Calibration did not validate a token contrast. The predeclared {selected_short}/{selected_adequate} caps are engineering diagnostic settings, not validated treatment thresholds."
    else:
        calibration_statement = "Calibration has not produced a valid contrast or a stored cap selection; no validated treatment thresholds are available."
    paragraphs = [
        ("Evidence and outcome", f"{evidence}\n\nRun status: {phase}. {outcome}\n\n{calibration_statement}" + (f"\n\nRecorded stop reason: {stop}" if stop else "")),
        ("Model and intervention", f"Requested model: {summary['requested_model']}. Returned models: {', '.join(returned_models) or 'none recorded'}. Fixed reasoning effort: {summary['effort']}. Selected total-output caps: short {selected_short}, adequate {selected_adequate}. Exact prompts, state histories and responses are retained in manifest.json and requests/*.json. The allowance applies to total generated output, including reasoning tokens and the structured decision; nonreasoning tokens equal total output minus the reasoning subset. Unused allowance is not measured effort. Selected caps are recorded in selection.json / calibration_recommendation.json.\n\nThe common {timeout}-second transport timeout is operational reliability protection, not a deadline treatment. Recorded latency includes service and transport delay. Tokens do not have a validated conversion to human thinking seconds; the study manipulates an LLM computational allowance. A separate time-limit experiment requires separate calibration. Usage summaries include measured requests only; unobserved usage remains missing, and its reserved charge remains in the ledger."),
        ("Decisions, failures and costs", f"Zero is a permitted decision. No positive-order floor or advertised order maximum is imposed. Concerning zero orders (positive backlog with no stock or outstanding replenishment) are reported without being changed or excluded. Fallback orders belong to a disclosed diagnostic continuation policy, not the model. A completed chain with fallbacks does not become a complete model-behavior sample. Full costs are descriptive whenever fallbacks occur; incomplete chains have only accumulated observed costs.\n\nCosts include all {game_weeks} game weeks, including {warmup_weeks} fixed warmup weeks, plus {continuation_weeks} automatic continuation weeks. Automatic and fallback orders are excluded from valid model decision zero rates. Raw failures and their charges/reservations remain in the request ledger; no failed generation is converted into a valid zero decision."),
        ("Bullwhip and beta", f"Bullwhip ratios use sample variance over the full {game_weeks} game weeks including warmup, or the observed partial window when stopped. Both local incoming demand and external customer demand are reported. A zero-variance denominator produces a missing ratio. The single demand step and different timing upstream make these descriptive ratios sensitive to the window; a ratio alone does not establish sustained oscillation.\n\nBeta is an exploratory fitted supply-line weight, not a direct measurement of thinking or a diagnosed human bias. A role with invalid/fallback decisions has an unusable fit. Variation, fit quality and identification flags remain visible. Chain beta requires four usable role fits. No confirmatory p-values are calculated for this small feasibility pilot; weekly calls are not independent participants."),
        ("Research scope", "The common persona is a hypothetical second-year business student at University at Buffalo. A demographic prompt does not validate a population model. Repeated calls represent model-generated decisions, not independent University at Buffalo students. Claims about actual students require answers from real students using comparable tasks; see Bisbee et al. (https://doi.org/10.1017/pan.2024.5) on the limits of synthetic survey respondents. The cited GPT behavioral-operations study (https://doi.org/10.1287/msom.2023.0279) motivates investigating model decisions across operational tasks; this Beer Game alone does not diagnose all eighteen biases. Deliberately requesting a bias or oscillation would change the behavior being studied. Prompt, model, effort, budget selection and failure policy must remain fixed within the comparison, and historical none-effort protocols must remain separate."),
        ("Sources and reproducibility", "Game timing/cost background: MIT / John Sterman Beer Game (https://web.mit.edu/jsterman/www/SDG/beergame.html). API token accounting and incomplete responses: OpenAI reasoning guide (https://developers.openai.com/api/docs/guides/reasoning). Response schema mechanics: OpenAI structured outputs (https://developers.openai.com/api/docs/guides/structured-outputs). These sources support game/API mechanics; they do not validate student realism or tokens as human time.\n\nThe manifest, raw requests, calibration, rule checks and trajectory CSVs are the reproducible evidence. Currency values are local estimated charges/reservations, not an account balance; unresolved request usage may remain reserved. Software checks and offline synthetic runs cannot supply missing live results."),
    ]
    tables = {
        "Condition outcomes": conditions, "Group outcomes": groups, "Role metrics": roles,
        "Beta diagnostics": fits, "Calibration": calibration_summary, "Calibration states": calibration, "Rules checks": rules_frame,
        "All game rounds": rounds, "Automatic continuation": tail, "Request index": requests,
        "Configuration": pd.DataFrame([dict(setting=k, value=json.dumps(v) if isinstance(v, (dict, list)) else v) for k, v in config.items()]),
    }
    for name, frame in (("group_outcomes", groups), ("role_metrics", roles), ("beta_fits", fits), ("condition_outcomes", conditions), ("request_index", requests), ("calibration_report_summary", calibration_summary)):
        frame.to_csv(directory / f"{name}.csv", index=False)
    warnings, plots = [], []
    try:
        plots = _plots(directory, rounds, tail, requests, synthetic)
    except ImportError as exc:
        warnings.append(f"Plot export unavailable: {exc.name} dependency missing.")
    try:
        _workbook(directory, tables, summary)
    except ImportError as exc:
        warnings.append(f"Excel export unavailable: {exc.name} dependency missing.")
    try:
        _docx(directory, paragraphs, tables, plots)
    except ImportError as exc:
        warnings.append(f"Word export unavailable: {exc.name} dependency missing.")
    summary["artifact_warnings"] = warnings
    (directory / "report_summary.json").write_text(json.dumps(_finite(summary), indent=2, allow_nan=False) + "\n", encoding="utf-8")
    lines = ["# Student-persona Beer Game: evidence report", "", evidence, "", outcome,
             "", f"Estimated charges/reservations: ${summary['estimated_charged_or_reserved_usd']:.4f}; saved requests: {len(requests)}.",
             "", f"Saved failed decision requests: {summary['saved_failed_decision_requests']}; committed fallback orders: {fallback_count}. Stopped batches may contain responses that were never committed. Unknown/unmeasured usage requests: {summary['unresolved_or_unmeasured_usage_requests']}."]
    for title, text in paragraphs:
        lines += ["", "## " + title, "", text]
    for name, columns in (
        ("Condition outcomes", ["condition", "chains_complete", "clean_chains", "valid_submitted_decisions", "missing_decisions", "fallback_count", "zero_rate_valid", "max_submitted_order", "mean_total_scored_cost"]),
        ("Group outcomes", ["group_id", "full_game", "fallback_count", "total_scored_cost", "usable_role_betas", "group_beta"]),
        ("Calibration", None),
        ("Beta diagnostics", ["group_id", "role", "beta", "usable", "fit_flag"]),
    ):
        lines += ["", "## " + name, "", _table(tables[name], columns)]
    lines += ["", calibration_statement + " Full selection criteria and state results are saved in [calibration_recommendation.json](calibration_recommendation.json) and [calibration.csv](calibration.csv).",
              "", "Primary references: " + "; ".join(f"[{title}]({url})" for title, url in SOURCES) + "."]
    for name in plots:
        lines += ["", f"![{name.replace('_', ' ').replace('.png', '')}]({name})"]
    if warnings:
        lines += ["", "## Export availability", "", "\n".join(warnings)]
    (directory / "REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return directory
