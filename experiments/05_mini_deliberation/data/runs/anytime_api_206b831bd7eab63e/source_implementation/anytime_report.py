"""Saved-evidence analysis and report export. No API client or calls."""
from __future__ import annotations

import hashlib
import json
import math
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

from .game import BeerGame, ROLES
from .estimation import estimate_beta
from .reasoning_protocol import decide, digest, save
from .anytime_protocol import display_state, demand_schedule, TASK
from .student_profiles import assignment


def read_csv(path):
    try:
        return pd.read_csv(path)
    except (FileNotFoundError, pd.errors.EmptyDataError):
        return pd.DataFrame()


def variance_ratio(numerator, denominator):
    variance = float(np.var(denominator, ddof=1)) if len(denominator) >= 2 else 0
    return float(np.var(numerator, ddof=1)/variance) if variance > 0 else None


def paired_inference(differences, eligible):
    """Paired block inference, never weekly pseudoreplication.

    A t interval/test assumes independent blocks and approximately normal paired
    differences. Eight is only a reporting floor, not a power calculation.
    A degenerate variance is flagged rather than manufactured into significance.
    """
    d = np.asarray(differences, dtype=float)
    result = dict(pairs=len(d), estimate=float(np.mean(d)) if len(d) else None,
                  p_raw=None, ci_lower=None, ci_upper=None, status='descriptive; inference withheld')
    if eligible and len(d) >= 8 and np.std(d, ddof=1) > 0:
        se = float(stats.sem(d))
        interval = float(stats.t.ppf(.975, len(d)-1))*se
        result.update(p_raw=float(stats.ttest_1samp(d, 0).pvalue),
                      ci_lower=result['estimate']-interval, ci_upper=result['estimate']+interval,
                      status='exploratory paired t test; inspect block differences and assumptions')
    return result


def holm_family(rows):
    ordered = sorted(enumerate(rows), key=lambda item: item[1]['p_raw'] if item[1]['p_raw'] is not None else 1)
    running = 0.0
    for rank, (_, row) in enumerate(ordered):
        if row['p_raw'] is None:
            row['p_holm'] = None
        else:
            running = max(running, min(1., (len(rows)-rank)*row['p_raw']))
            row['p_holm'] = running
    return rows


def planned_pairs(standardized_paired_effect, power=.8, family_alpha=.05, tests=3):
    """Prospective planning under normal paired differences, not promised power."""
    if not math.isfinite(standardized_paired_effect) or standardized_paired_effect <= 0:
        raise ValueError('Specify a positive meaningful effect divided by the SD of paired chain differences.')
    for n in range(3, 10001):
        critical = stats.t.ppf(1-family_alpha/(2*tests), n-1)
        ncp = standardized_paired_effect*np.sqrt(n)
        achieved = stats.nct.sf(critical, n-1, ncp)+stats.nct.cdf(-critical, n-1, ncp)
        if achieved >= power:
            return dict(pairs=n, standardized_paired_effect=standardized_paired_effect,
                        assumed_power=power, achieved_under_assumptions=float(achieved),
                        assumptions='Independent normal paired differences; Bonferroni planning for three two-sided tests. Requires external effect/variance justification.')
    raise ValueError('Required sample exceeds 10,000 pairs.')


def audit_evidence(directory, manifest, rounds, probes, records):
    by_key = {r['request_key']: r for r in records if r.get('status') == 'ok'}
    for name, expected in manifest['source_hashes'].items():
        if digest((directory/'source_implementation'/name).read_text(encoding='utf-8')) != expected:
            raise ValueError('Saved scientific source hash mismatch: '+name)
    for record in records:
        fields = ('protocol', 'identity', 'messages', 'limit', 'schema', 'requested_model',
                  'reasoning', 'mode', 'config_fingerprint')
        if digest({k: record.get(k) for k in fields}) != record['request_key']:
            raise ValueError('Request fingerprint mismatch.')
        if (record['protocol'] != manifest['protocol'] or record['reasoning'] != {'effort': 'high'}
                or record['requested_model'] != manifest['config']['model']
                or record['config_fingerprint'] != manifest['fingerprint']):
            raise ValueError('Request model/protocol/effort mismatch.')
    for frame in (rounds[rounds.phase == 'model'] if not rounds.empty else rounds, probes):
        for row in frame.to_dict('records'):
            selected = by_key.get(row['selected_request_key'])
            if selected is None or not decide(selected).valid or decide(selected).order != row['order']:
                raise ValueError('Scored choice is not an exact valid model submission.')
            if row['initial_request_key'] not in by_key:
                raise ValueError('Initial decision provenance missing.')
            first = by_key[row['initial_request_key']]
            initial = decide(first)
            if not initial.valid or initial.order != row['initial_order']:
                raise ValueError('Initial choice mismatch.')
            if row['order_source'] == 'revision' and row['selected_request_key'] != row['revision_request_key']:
                raise ValueError('Revision provenance mismatch.')
            revised = by_key.get(row.get('revision_request_key'))
            total = sum((r.get('usage') or {}).get('output_tokens', 0) for r in (first, revised) if r)
            if total != row['total_output_tokens']:
                raise ValueError('Decision token totals mismatch.')
    count = 0
    if not rounds.empty:
        for (block, condition), chain in rounds.groupby(['block', 'condition']):
            game = BeerGame()
            histories = {role: [] for role in ROLES}
            path = demand_schedule(manifest['config']['weeks'], manifest['config']['seed']+10000+int(block))
            assigned = assignment(manifest['config']['seed'], int(block))
            for week, own in chain.groupby('week', sort=True):
                if int(week) != game.week+1 or set(own.role) != set(ROLES):
                    raise ValueError('Noncontiguous or partial committed week.')
                ordered = own.set_index('role').loc[list(ROLES)]
                replay = game.observe(int(ordered.iloc[0].customer_demand))
                if int(ordered.iloc[0].customer_demand) != path[min(int(week)-1, len(path)-1)]:
                    raise ValueError('Demand differs from the prospective exogenous path.')
                for i, (state, (_, row)) in enumerate(zip(replay, ordered.iterrows())):
                    for field in ('inventory', 'backlog', 'received', 'shipped', 'demand', 'supply_line', 'cost'):
                        if state[field] != row[field]:
                            raise ValueError('Physical replay mismatch: '+field)
                    if row.profile_id != assigned[i]:
                        raise ValueError('Persona-role assignment mismatch.')
                    shown = display_state(game, state) | {'profile_id': assigned[i]}
                    if row.phase == 'model':
                        first = by_key[row.initial_request_key]
                        payload = dict(task=TASK, participant_profile=manifest['student_profiles'][assigned[i]],
                                       past_rounds=histories[state['role']],
                                       current_state={k: v for k, v in shown.items() if k != 'profile_id'})
                        expected = [{'role': 'system', 'content': manifest['common_prompt']},
                                    {'role': 'user', 'content': json.dumps(payload, separators=(',', ':'), allow_nan=False)}]
                        if first['messages'] != expected:
                            raise ValueError('Initial prompt does not match the role, persona, state and history.')
                        ident = first['identity']
                        if ident.get('block') != block or ident.get('week') != week or ident.get('role') != state['role']:
                            raise ValueError('Initial request reused across independent decisions.')
                        revised = by_key.get(row.revision_request_key)
                        if condition == 'adequate':
                            if revised is None:
                                raise ValueError('Missing assigned revision opportunity.')
                            payload.update(submitted_order=int(row.initial_order),
                                           opportunity='You have one opportunity to reconsider. Keep or change your submitted order.')
                            expected[1]['content'] = json.dumps(payload, separators=(',', ':'), allow_nan=False)
                            if revised['messages'] != expected:
                                raise ValueError('Revision prompt mismatch.')
                        elif revised is not None:
                            raise ValueError('Short condition received an extra computation opportunity.')
                    if row.phase != 'automatic':
                        histories[state['role']].append(shown | {'order': int(row.order)})
                    count += 1
                game.commit([int(q) for q in ordered.order])
    return dict(passed=True, physical_states_replayed=count, request_records=len(records),
                selected_orders_verified=(int((rounds.phase == 'model').sum()) if not rounds.empty else 0)+len(probes))


def analyze(directory):
    directory = Path(directory)
    manifest = json.loads((directory/'manifest.json').read_text(encoding='utf-8'))
    progress = json.loads((directory/'progress.json').read_text(encoding='utf-8'))
    config = manifest['config']
    rounds, probes = read_csv(directory/'rounds.csv'), read_csv(directory/'probes.csv')
    calibration = read_csv(directory/'calibration.csv')
    screen_path = directory/'manipulation.json'
    screen = json.loads(screen_path.read_text(encoding='utf-8')) if screen_path.exists() else {'passed': False}
    records = [json.loads(p.read_text(encoding='utf-8')) for p in sorted((directory/'requests').glob('*.json'))]
    audit = audit_evidence(directory, manifest, rounds, probes, records)
    groups, role_metrics, contrasts, beta_fits = [], [], [], []
    if not rounds.empty:
        for (block, condition), chain in rounds.groupby(['block', 'condition']):
            decisions = chain[chain.phase == 'model']
            complete = len(chain) == (config['weeks']+12)*4
            ratios = []
            for role, role_rows in decisions.groupby('role'):
                beta_fits.append(dict(block=int(block), condition=condition, role=role, **estimate_beta(role_rows)))
                ratio = variance_ratio(role_rows.order, role_rows.customer_demand)
                ratios.append(ratio)
                role_metrics.append(dict(block=int(block), condition=condition, role=role,
                    profile_id=role_rows.iloc[0].profile_id,
                    bullwhip_customer=ratio, bullwhip_local=variance_ratio(role_rows.order, role_rows.demand),
                    demand_variance=float(role_rows.customer_demand.var()), order_variance=float(role_rows.order.var()),
                    zero_orders=int((role_rows.order == 0).sum()), max_order=int(role_rows.order.max()),
                    backlog_max=int(role_rows.backlog.max())))
            sensitivity = None
            if not probes.empty:
                own = probes[(probes.block == block) & (probes.condition == condition)]
                if len(own) == 8:
                    wide = own.pivot(index='role', columns='level', values='order')
                    sensitivity = float(((wide.low-wide.high)/24).mean())
            groups.append(dict(block=int(block), condition=condition, complete=complete,
                               scored_cost=float(chain.cost.sum()) if complete else None,
                               observed_cost=float(chain.cost.sum()),
                               decision_cost=float(decisions.cost.sum()),
                               bullwhip=float(np.mean(ratios)) if complete and len(ratios) == 4 and all(r is not None for r in ratios) else None,
                               sensitivity=sensitivity, decisions=len(decisions),
                               mean_tokens=float(decisions.total_output_tokens.mean()) if len(decisions) else None,
                               median_tokens=float(decisions.total_output_tokens.median()) if len(decisions) else None))
    groups = pd.DataFrame(groups)
    token_ratio = None
    main_pass = False
    if not rounds.empty and 'total_output_tokens' in rounds:
        decisions = rounds[rounds.phase == 'model']
        medians = decisions.groupby('condition').total_output_tokens.median()
        if set(medians.index) == {'short', 'adequate'} and medians['adequate'] > 0:
            token_ratio = float(medians['short']/medians['adequate'])
            main_pass = token_ratio <= .75
    eligible = (config['mode'] == 'api' and screen['passed'] and main_pass
                and progress['phase'] == 'complete' and not groups.empty
                and len(groups) == 2*config['groups_per_condition'] and bool(groups.complete.all()))
    hypotheses = []
    for name, metric, direction in [('H1_sensitivity', 'sensitivity', '<0'), ('H2_cost', 'scored_cost', '>0'),
                                    ('H3_bullwhip', 'bullwhip', '>0')]:
        differences = []
        if not groups.empty:
            for block, pair in groups.groupby('block'):
                if set(pair.condition) == {'short', 'adequate'} and pair[metric].notna().all():
                    indexed = pair.set_index('condition')
                    delta = float(indexed.loc['short', metric]-indexed.loc['adequate', metric])
                    differences.append(delta)
                    contrasts.append(dict(hypothesis=name, block=int(block), short=float(indexed.loc['short', metric]),
                                          adequate=float(indexed.loc['adequate', metric]), difference=delta))
        hypotheses.append(dict(hypothesis=name, predicted_short_minus_adequate=direction,
                               **paired_inference(differences, eligible)))
    return dict(manifest=manifest, progress=progress, screen=screen, audit=audit, rounds=rounds,
                probes=probes, calibration=calibration, groups=groups, roles=pd.DataFrame(role_metrics),
                hypotheses=pd.DataFrame(holm_family(hypotheses)), contrasts=pd.DataFrame(contrasts), beta_fits=pd.DataFrame(beta_fits),
                actual_token_ratio=token_ratio, main_manipulation_passed=bool(main_pass),
                charge_usd=sum(r.get('charge_usd', 0) for r in records), records=records)


def markdown_table(frame):
    if frame.empty:
        return 'No observations available.'
    clean = frame.astype(object).where(pd.notna(frame), '—')
    def cell(value):
        return (f'{value:.6g}' if isinstance(value, float) else str(value)).replace('|', '/')
    return '\n'.join(['| '+' | '.join(frame.columns)+' |', '| '+' | '.join(['---']*len(frame.columns))+' |']+
                     ['| '+' | '.join(map(cell, row))+' |' for row in clean.itertuples(index=False, name=None)])


def build_report(directory, root=None):
    directory = Path(directory).resolve()
    root = Path(root or Path(__file__).resolve().parents[2])
    result = analyze(directory)
    manifest, progress = result['manifest'], result['progress']
    config = manifest['config']
    output = root/'reports'/directory.name
    inputs = {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in sorted(directory.rglob('*')) if p.is_file() and p.suffix in ('.json', '.csv', '.py')}
    signature = digest(dict(inputs=inputs, analysis=Path(__file__).read_text(encoding='utf-8')))
    old_path = output/'analysis_manifest.json'
    if old_path.exists() and json.loads(old_path.read_text(encoding='utf-8'))['analysis_fingerprint'] != signature:
        old = json.loads(old_path.read_text(encoding='utf-8'))['analysis_fingerprint']
        archive = root/'reports/archive/analysis_versions'/directory.name/old
        if not archive.exists():
            shutil.copytree(output, archive)
    output.mkdir(parents=True, exist_ok=True)
    synthetic = config['mode'] == 'offline'
    title = 'SYNTHETIC SOFTWARE CHECK — no model or human observations' if synthetic else 'LLM Beer Game: initial choice and additional deliberation'
    tables = {k: result[k] for k in ('hypotheses', 'groups', 'roles', 'contrasts', 'calibration', 'probes', 'beta_fits')}
    tables['student_profiles'] = pd.DataFrame(manifest['student_profiles'].values())
    sections = [f'# {title}',
        f"Run `{directory.name}`. Protocol `{manifest['protocol']}`. Status: **{progress['phase']}**. "
        f"Reason: {progress.get('reason', 'Finished requested stage.')} API charges/reservations for this run: **${result['charge_usd']:.6f}**.",
        '## What this experiment tests',
        'A hypothetical second-year UB business student makes an initial current-week choice at fixed high effort. '
        'The short condition submits that choice. The adequate condition receives one additional high-effort revision opportunity. '
        'Initial generation has a 16,384-token ceiling; the extra opportunity has an 8,192-token ceiling. '
        'This is an additional-deliberation intervention, not a single-call total-budget intervention or a validated human deadline. '
        'Information quantity is held constant; information overload is not separately randomized. '
        'Four fictional personality profiles are matched across treatments and rotated across roles in four-pair blocks. '
        'Their equal weighting is a design choice, not a measured student population. No human data exist.',
        'Only valid same-week model submissions become orders. An unfinished token-limited revision leaves the initial submission in force. '
        'An initial failure or technical error stops the run. Zero and large orders remain outcomes; no interior-order target is imposed.',
        '## Manipulation checks',
        f"Held-out and development feasibility screen passed: **{result['screen']['passed']}**. "
        f"Main-game median short/adequate total-token ratio: **{result['actual_token_ratio']}**; required ≤0.75. "
        f"Main-game separation passed: **{result['main_manipulation_passed']}**. "
        'These are small-sample computational checks, not population validation. Both phases count toward computation. '
        'No games are purchased after a failed live feasibility screen.',
        markdown_table(pd.DataFrame(result['screen'].get('splits', []))),
        '## Hypotheses and conclusions', markdown_table(result['hypotheses']),
        'H3 is the primary bullwhip outcome: mean over roles of sample order variance divided by customer-demand variance during '
        'model decision weeks only. H2 is total chain holding/backlog cost including the fixed practice and automatic continuation. '
        'H1 is a separate matched-state mechanism assay: negative order change divided by the 24-case pipeline increase, averaged across roles. '
        'H1 measures direct supply-line responsiveness; it does not identify the old dynamic beta. The historical beta hypothesis remains separate.',
        'All contrasts are short minus adequate. Independent replication is the paired four-role chain block, not a call or week. '
        'Initial draws are shared only within a pair when complete prompts match, a common-random-draw variance-reduction design. '
        'Different blocks always buy distinct draws. Paired t tests and 95% intervals require independent blocks and suitable difference distributions. '
        'The three tests form one Holm family. Fewer than eight pairs, failed manipulation, incomplete games or synthetic inputs withhold inference. '
        'Eight is a reporting floor, not evidence of adequate power. No sample size guarantees significance.',
        '## Chain outcomes', markdown_table(result['groups']),
        '## Bullwhip and behavior by role', markdown_table(result['roles']),
        'Customer demand is prospectively generated as independent uniform integers 0 through 8 after four practice weeks. '
        'Participants know the distribution; this reduces a demand-learning confound. '
        'Each paired block receives the same exogenous sequence, hidden from participants. This preserves variation in the measurement window. '
        'A local ratio with constant local demand remains undefined. Twelve automatic weeks repeat each role\'s final submitted quantity '
        'and the final customer demand; this differs from the historical demand-following tail and is not behavioral data.',
        '## Cost controls and reproducibility',
        'Identical paired initial prompts share one paid draw; repeated independent replications never share draws. '
        'Exact completed request checkpoints are reused on resume. Stable instructions occur first for provider prompt caching; '
        'actual cached-input usage is recorded and output/reasoning charges remain payable. Temperature is omitted in all calls. '
        'No correctness retry or automatic switch to a more expensive model is used. '
        'The $10 campaign includes the previous $5.19360675; new campaign identifiers require new explicit budget authorization.',
        f"Saved-evidence audit: **{result['audit']}**. Exact requests, responses, chosen-source request IDs, code snapshots, "
        'demand paths and configuration are retained under the source data directory. Analysis makes no API calls.',
        '## Limitations and next decision',
        'A completed pilot establishes operation and, if the checks pass, a measured extra-computation contrast. '
        'It does not establish that an LLM is a real UB student, that information overload was manipulated, or that extra computation must improve decisions. '
        'Keep null or opposite-direction outcomes. Determine a meaningful chain-level effect and power before a larger study. '
        'Human comparison materials are available in docs/STUDENT_VALIDATION_20261011.md.',
        '## Calibration observations', markdown_table(result['calibration']),
        '## Student profiles', markdown_table(tables['student_profiles']),
        'Profiles cross an organized/methodical versus flexible/informal tendency with emotionally steady versus uncertainty-sensitive framing. '
        'These are explicit fictional personality prompts, not psychometric scores or instructions to exhibit a bias. '
        'See reports/revised_design_20261011/REPORT.md for the primary literature and the limits of this adaptation.',
        '## Original structural beta: exploratory identification check', markdown_table(result['beta_fits']),
        'The original dynamic beta fit is retained with its original identification checks. Unusable numerical values do not measure supply-line neglect. '
        'The H1_sensitivity assay above is a revised operationalization of the mechanism, not a relabeling of these coefficients.',
        '## Primary API documentation',
        '[Reasoning output limits](https://developers.openai.com/api/docs/guides/reasoning), '
        '[model and pricing](https://developers.openai.com/api/docs/models/gpt-5.4-mini), '
        '[prompt caching](https://developers.openai.com/api/docs/guides/prompt-caching).']
    report = '\n\n'.join(sections)+'\n'
    (output/'REPORT.md').write_text(report, encoding='utf-8')
    (output/'PROMPT.txt').write_text(manifest['common_prompt'], encoding='utf-8')
    for name, table in tables.items():
        table.to_csv(output/f'{name}.csv', index=False)
    with pd.ExcelWriter(output/'DETAILS.xlsx') as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
    from docx import Document
    document = Document()
    document.add_heading(title, 0)
    for section in sections[1:]:
        if section.startswith('## '):
            document.add_heading(section[3:], level=1)
        elif section.startswith('| '):
            lines = section.splitlines()
            table = document.add_table(rows=1, cols=len(lines[0].strip('|').split('|')))
            table.style = 'Light Shading Accent 1'
            for cell, value in zip(table.rows[0].cells, lines[0].strip('|').split('|')):
                cell.text = value.strip()
            for line in lines[2:]:
                for cell, value in zip(table.add_row().cells, line.strip('|').split('|')):
                    cell.text = value.strip()
        else:
            document.add_paragraph(section.replace('**', '').replace('`', ''))
    document.save(output/'Anytime_Study_Report.docx')
    if not result['rounds'].empty:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        figure, axes = plt.subplots(2, 2, figsize=(11, 7), sharex=True)
        for role, axis in zip(ROLES, axes.flat):
            frame = result['rounds'][(result['rounds'].role == role) & (result['rounds'].phase != 'automatic')]
            for (block, condition), own in frame.groupby(['block', 'condition']):
                axis.plot(own.week, own.order, label=f'{condition}, pair {block}')
            demand = frame[frame.condition == 'short'].groupby('week').customer_demand.mean()
            axis.plot(demand.index, demand, ':', color='black', label='customer demand')
            axis.set_title(role); axis.set_ylabel('Cases'); axis.set_xlabel('Week'); axis.legend(fontsize=7)
        figure.suptitle(title); figure.tight_layout(); figure.savefig(output/'orders.png', dpi=160); plt.close(figure)
    save(output/'analysis_manifest.json', dict(analysis_fingerprint=signature, source_run=str(directory),
         inputs_sha256=inputs, audit=result['audit'], additional_api_charges=0, synthetic=synthetic))
    pointer = dict(run=str(directory), report=str(output.resolve()), analysis_fingerprint=signature, synthetic=synthetic)
    save(root/'reports/LATEST_ANYTIME_REPORT.json', pointer)
    if not synthetic:
        (root/'reports/LATEST_REPORT.txt').write_text('\n'.join(str((output/name).resolve()) for name in
              ('Anytime_Study_Report.docx', 'REPORT.md', 'DETAILS.xlsx'))+'\n', encoding='utf-8')
    return output
