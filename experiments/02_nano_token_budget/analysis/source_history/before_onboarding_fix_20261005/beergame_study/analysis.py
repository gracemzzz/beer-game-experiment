import json,itertools,copy
from pathlib import Path
import numpy as np
import pandas as pd
from .game import BeerGame,ROLES,demand_path,continuation
from .api import save
from .estimation import estimate_beta,predicted_orders


def ratio(numerator,denominator):
    return float(np.var(numerator,ddof=1)/np.var(denominator,ddof=1)) if len(denominator)>1 and np.var(denominator,ddof=1)>1e-12 else np.nan

def signflip(differences,seed=7):
    d=np.array(differences,float);n=len(d)
    if n<=16:
        values=np.abs(np.array(list(itertools.product((-1,1),repeat=n)))@d/n)
        return float(np.mean(values>=abs(d.mean())-1e-12))
    rng=np.random.default_rng(seed)
    values=np.abs(rng.choice([-1,1],size=(50000,n))@d/n)
    return float((1+np.sum(values>=abs(d.mean())-1e-12))/50001)

def hypothesis_table(groups,config,complete,replaced):
    records=[]
    eligible=(config['mode']=='api' and complete and not replaced and config['groups_per_condition']>=8)
    for hypothesis,outcome,direction in [('H1','group_beta','short minus adequate < 0'),('H2','total_scored_cost','short minus adequate > 0')]:
        pivot=groups.pivot(index='block',columns='condition',values=outcome).reindex(columns=['short','adequate'])
        all_usable=len(pivot)==config['groups_per_condition'] and pivot.notna().all().all()
        d=(pivot['short']-pivot['adequate']) if all_usable else pd.Series(dtype=float)
        records.append(dict(hypothesis=hypothesis,outcome=outcome,estimate=float(d.mean()) if len(d) else np.nan,
            complete_pairs=int(pivot.notna().all(axis=1).sum()),predicted_direction=direction,
            p_raw=signflip(d) if eligible and all_usable else np.nan,
            status='exploratory paired group test' if eligible and all_usable else
                'withheld: offline, fewer than 8 pairs, incomplete sample, replacement, or missing outcome'))
    # Preserve the planned family of two tests even if H1 is unidentifiable.
    ps=np.array([r['p_raw'] if np.isfinite(r['p_raw']) else 1. for r in records]);idx=np.argsort(ps)
    adjusted=np.empty(2);adjusted[idx[0]]=min(1.,2*ps[idx[0]]);adjusted[idx[1]]=max(adjusted[idx[0]],ps[idx[1]])
    for i,r in enumerate(records):r['p_holm']=adjusted[i] if np.isfinite(r['p_raw']) else np.nan
    return pd.DataFrame(records)


def game_tables(frame,tail,config,complete,replaced):
    roles=[];fits=[];groups=[]
    for gid,g in frame.groupby('group_id'):
        customer=g[g.role=='retailer'].sort_values('week').demand.to_numpy()
        for role,a in g.groupby('role'):
            a=a.sort_values('week');agent=a[a.phase=='agent'];meta=a.iloc[0]
            fit=estimate_beta(agent)
            # The eight-interior-order cap rule is not used: there is no order-entry cap.
            # Identification still requires rank/variation checks and meaningful positive decisions.
            if int((agent.order>0).sum())<5:
                fit.update(usable=False,fit_flag=fit['fit_flag']+';too_few_positive_orders')
            fits.append(dict(group_id=gid,role=role,condition=meta.condition,block=int(meta.block),
                positive_orders=int((agent.order>0).sum()),**fit))
            roles.append(dict(group_id=gid,role=role,condition=meta.condition,
                holding_cost=float((.5*a.inventory).sum()),backlog_cost=float(a.backlog.sum()),
                bullwhip_vs_local=ratio(a.order.to_numpy(),a.demand.to_numpy()),
                bullwhip_vs_customer=ratio(a.order.to_numpy(),customer),
                order_variance=float(a.order.var(ddof=1)),mean_order=float(agent.order.mean()),
                zero_rate=float((agent.order==0).mean()),max_order=int(agent.order.max()),
                backlog_weeks=int((a.backlog>0).sum()),max_backlog=int(a.backlog.max()),
                terminal_backlog=int(a.backlog.iloc[-1])))
        retailer=g[g.role=='retailer'];t=tail[tail.group_id==gid];meta=g.iloc[0]
        group_fits=[fit for fit in fits if fit['group_id']==gid]
        groups.append(dict(group_id=gid,block=int(meta.block),condition=meta.condition,
            agent_phase_cost=float(g.cost.sum()),continuation_cost=float(t.cost.sum()),
            total_scored_cost=float(g.cost.sum()+t.cost.sum()),
            group_beta=float(np.mean([x['beta'] for x in group_fits])) if all(x['usable'] for x in group_fits) else np.nan,
            usable_role_betas=sum(x['usable'] for x in group_fits),
            customer_fulfillment=float(retailer.shipped.sum()/retailer.demand.sum()),
            retailer_backlog_weeks=int((retailer.backlog>0).sum()),
            terminal_inventory=int(t[t.week==t.week.max()].inventory.sum()),
            terminal_backlog=int(t[t.week==t.week.max()].backlog.sum())))
    groups=pd.DataFrame(groups)
    return pd.DataFrame(roles),pd.DataFrame(fits),groups,hypothesis_table(groups,config,complete,replaced)


def software_checks():
    tests=[]
    costs={}
    for policy in ('zero','demand','large'):
        game=BeerGame();total=0
        for week,demand in enumerate(demand_path(40),1):
            states=game.observe(demand);total+=sum(s['cost'] for s in states)
            orders=[4]*4 if week<=4 else [0 if policy=='zero' else (40 if policy=='large' else s['demand']) for s in states]
            game.commit(orders)
        total+=sum(s['cost'] for s in continuation(game,demand))
        costs[policy]=total
    tests.append(dict(check='Simple demand benchmark beats always-zero and always-40',passed=costs['demand']<min(costs['zero'],costs['large']),detail=str(costs)))
    game=BeerGame()
    for week,d in enumerate(demand_path(40),1):
        state=game.observe(d)
        if week<40:game.commit([s['demand'] for s in state])
    for role in range(4):
        totals=[]
        for order in (0,8,40):
            trial=copy.deepcopy(game);orders=[8]*4;orders[role]=order;trial.commit(orders)
            totals.append(sum(s['cost'] for s in continuation(trial,8)))
        tests.append(dict(check='Final order affects score: '+ROLES[role],passed=len(set(totals))>1,detail=str(totals)))
    rng=np.random.default_rng(6);n=120
    d=rng.integers(2,15,n);ni=rng.normal(10,15,n);sl=rng.uniform(0,60,n)
    y=predicted_orders([.35,.4,.65,25],d,ni,sl)
    fit=estimate_beta(pd.DataFrame(dict(week=np.arange(n),order=y,demand=d,net_inventory=ni,supply_line=sl)))
    tests.append(dict(check='Known beta recovery, including valid zero orders',passed=bool(fit['usable'] and abs(fit['beta']-.65)<.01),detail=str(fit['beta'])))
    tests.append(dict(check='Bullwhip ratio undefined for constant denominator',passed=bool(np.isnan(ratio([1,3,9],[4,4,4]))),detail='NaN, not infinity or zero'))
    result=pd.DataFrame(tests);assert result.passed.all(),result.to_string();return result


def report(directory):
    directory=Path(directory);manifest=json.loads((directory/'manifest.json').read_text());config=manifest['config']
    status=json.loads((directory/'progress.json').read_text());synthetic=config['mode']=='offline'
    records=[json.loads(p.read_text()) for p in (directory/'requests').glob('*.json')]
    charge=sum(r.get('charge_usd',0) for r in records)
    lines=['# Token-budget Beer Game — conclusions', '',
        '1. '+('SYNTHETIC SOFTWARE CHECK: no GPT evidence or hypothesis inference.' if synthetic else 'Real API responses; LLM simulation, not human evidence.'),
        f'2. Run status: {status["phase"]}. Estimated cost/reservations: ${charge:.4f}.',
        f'3. Saved successful API stages: {sum(r["status"]=="ok" for r in records)}; other saved attempts: {sum(r["status"]!="ok" for r in records)}.']
    if status.get('reason'):lines.append('Stop reason: '+status['reason'])
    tables={}
    if config['stage']=='calibration' and (directory/'calibration_recommendation.json').exists():
        rec=json.loads((directory/'calibration_recommendation.json').read_text())
        lines += [f'4. Suggested memo caps: {rec["short_tokens"]} / {rec["adequate_tokens"]} tokens; doubled benchmark {rec["benchmark_tokens"]}.',
            f'5. Manipulation screen passed: {rec["manipulation_screen_passed"]}. This is a provisional feasibility check, not proof of a human time-pressure equivalent.',
            '6. '+('Review the state-level choices; then run one group per condition before scaling.' if rec['manipulation_screen_passed'] else
                   'Do not start a paid main experiment yet. Inspect truncation, actual token use and adequate-versus-benchmark agreement.'),
            'Calibration uses 24 independent hypothetical states (all four roles), not complete game chains. ' +
            ('P25/P95 refer to memo-token use under the 1024-token reference. These percentiles cannot directly represent human response times.'
             if config.get('calibration_method','percentile')=='percentile' else
             'This follow-up uses prospectively fixed caps and retains the earlier completed reference screen. It does not reselect percentiles or change the pass thresholds.')]
        tables['Calibration summary']=pd.read_csv(directory/'calibration_summary.csv')
        checks=tables['Calibration summary'].set_index('screen')
        lines += [f'Short truncation: {checks.loc["short","truncated_fraction"]:.1%} (required at least 20%); '
                  f'adequate truncation: {checks.loc["adequate","truncated_fraction"]:.1%} (required at most 5%).',
                  f'Adequate versus doubled allowance: median absolute order difference {rec["median_absolute_adequate_benchmark_order_difference"]:.1f} units (screen requires at most 4).']
        tables['Calibration state results']=pd.read_csv(directory/'calibration.csv')
    elif config['stage']=='game' and (directory/'rounds.csv').exists():
        frame=pd.read_csv(directory/'rounds.csv');tail=pd.read_csv(directory/'continuation.csv')
        ledger=json.loads((directory/'group_attempts.json').read_text())
        replaced=any(g['replacement']>0 for g in ledger)
        complete=status['phase']=='complete' and frame.group_id.nunique()==2*config['groups_per_condition']
        roles,fits,groups,hypotheses=game_tables(frame,tail,config,complete,replaced)
        for name,df in [('role_metrics',roles),('beta_fits',fits),('group_outcomes',groups),('hypotheses',hypotheses)]:
            df.to_csv(directory/(name+'.csv'),index=False)
        agent=frame[frame.phase=='agent']
        tokens=agent.groupby('condition').agg(mean_preparation_tokens=('plan_tokens','mean'),
            median_preparation_tokens=('plan_tokens','median'),truncated_fraction=('plan_truncated','mean'),
            mean_final_tokens=('final_tokens','mean'),zero_rate=('order',lambda x:float((x==0).mean())),
            max_order=('order','max')).reset_index()
        tokens.to_csv(directory/'manipulation_check.csv',index=False)
        lines += [f'4. Completed groups: {frame.group_id.nunique()}/{2*config["groups_per_condition"]}; usable beta: {int(fits.usable.sum())}/{len(fits)} roles.',
            '5. H1 tests lower beta under the short allowance; H2 tests higher total scored cost. H2 does not require usable beta. See the separate status for each test.',
            '6. Review token separation, failures, cost and bullwhip plots before increasing sample size. Zero orders and large orders are outcomes, not automatic grounds for discarding a group.',
            'Cost is holding plus backlog, including warmup and the 12 automatic terminal weeks. '
            'Automatic decisions never enter beta, token-use or bullwhip estimates. '
            'Fulfillment counts late delivery; it is not an on-time service rate. '
            'Bullwhip ratios use all game weeks including warmup because retailer demand is constant after the step; '
            'a constant denominator is reported as missing. The deterministic demand step limits generalization.',
            'Beta is a structural-model diagnostic, not profit. No artificial upper order cap is fitted. '
            'Rank, response strength, variation and positive-order checks flag unreliable fits. '
            'A sample of at least eight paired chains is a computational inference floor, not a power calculation. '
            'The reflection proposes ten chains per condition; no sample size guarantees significance.']
        tables.update({'Group outcomes':groups,'Token manipulation':tokens,'Hypothesis tests':hypotheses,
                       'Role performance and bullwhip':roles,'Beta diagnostics':fits})
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        fig,axes=plt.subplots(2,2,figsize=(12,7),sharex=True)
        for ax,role in zip(axes.flat,ROLES):
            for (cond,block),a in frame[frame.role==role].groupby(['condition','block']):
                ax.plot(a.week,a.order,label=f'{cond} chain {block}',alpha=.8)
            ax.set_title(role);ax.set_ylabel('Order units');ax.legend(fontsize=7)
        fig.suptitle('Synthetic software trajectories' if synthetic else 'GPT order trajectories (bullwhip diagnosis)')
        fig.tight_layout();fig.savefig(directory/'order_trajectories.png',dpi=150);plt.close(fig)
        fig,ax=plt.subplots(figsize=(9,4))
        for g,a in roles.groupby('group_id'):
            ax.plot(ROLES,a.set_index('role').reindex(ROLES).bullwhip_vs_customer,marker='o',label=g)
        ax.set_ylabel('Order variance / external-demand variance');ax.legend(fontsize=7)
        fig.tight_layout();fig.savefig(directory/'bullwhip.png',dpi=150);plt.close(fig)
        lines+=['','![Order trajectories](order_trajectories.png)','![Bullwhip by role](bullwhip.png)']
    else:lines.append('4. No completed game data yet. Partial rounds and individual requests remain saved; no invented final results.')
    if (directory/'debriefs.json').exists():
        tables['Optional model debriefs']=pd.DataFrame(json.loads((directory/'debriefs.json').read_text()))
        lines += ['', 'Optional debriefs are retrospective model self-reports, not causal evidence of internal reasoning.']
    for title,df in tables.items():
        lines+=['','## '+title,'','```',df.round(4).to_string(index=False),'```']
    lines += ['', '## Interpretation limits', '',
        'This protocol restricts visible preparation tokens plus a common final-response allowance, with reasoning effort none. '
        'It does not directly limit hidden cognitive effort, wall-clock time or human thinking time. '
        'Prompts cannot reliably count remaining tokens; the API hard cap is the enforced constraint. '
        'A truncated memo is accepted as a deadline outcome, then a common 192-token stage solicits the order. '
        'The final stage can still involve implicit computation, and unused token allowance is not effort used.',
        'Technical retries retain the same prompt and cap. Failed final responses can replace a whole group only within the configured limit; '
        'all failed attempts remain recorded. Replacement selection withholds inferential p-values. '
        'No order is clipped or replaced by a fabricated answer. The separate emergency scale check stops the run rather than making a boundary order.',
        '', 'Sources: [provided reflection](../../source/ThoughtsReflection_Paragraph.docx); '
        '[MIT Beer Game](https://web.mit.edu/jsterman/www/SDG/beergame.html); '
        '[Lee, Padmanabhan & Whang](https://pubsonline.informs.org/doi/10.1287/mnsc.43.4.546); '
        '[OpenAI model documentation](https://developers.openai.com/api/docs/models/gpt-5.4-nano).']
    text='\n'.join(lines)+'\n';(directory/'REPORT.md').write_text(text)
    from docx import Document
    from docx.shared import Pt
    doc=Document();doc.styles['Normal'].font.size=Pt(10)
    for line in lines:
        if line.startswith('# '):doc.add_heading(line[2:],0)
        elif line.startswith('## '):doc.add_heading(line[3:],1)
        elif line and line!='```':doc.add_paragraph(line)
    doc.save(directory/'REPORT.docx')
    with pd.ExcelWriter(directory/'DETAILS.xlsx',engine='openpyxl') as writer:
        pd.DataFrame([dict(phase=status['phase'],mode=config['mode'],estimated_cost=charge)]).to_excel(writer,sheet_name='Status',index=False)
        for title,df in tables.items():df.to_excel(writer,sheet_name=title[:31],index=False)
        for name in ('rounds','continuation','all_candidate_rounds'):
            if (directory/(name+'.csv')).exists():pd.read_csv(directory/(name+'.csv')).to_excel(writer,sheet_name=name[:31],index=False)
    return text
