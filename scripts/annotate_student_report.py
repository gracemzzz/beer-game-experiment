"""Add operational context to derived reports without changing scientific evidence."""
import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'src'))
from beergame_study.operational_io import install_windows_file_retry
install_windows_file_retry()
from beergame_study.reasoning_protocol import save


def annotate(directory, note_path):
    directory = Path(directory)
    summary_path = directory/'report_summary.json'
    original = summary_path.read_bytes()
    summary = json.loads(original)
    imported = json.loads((directory/'calibration_import.json').read_text(encoding='utf-8'))
    imported_charge = sum(json.loads((directory/'requests'/name).read_text(encoding='utf-8')).get('charge_usd', 0)
                          for name in imported['imported_records'])
    ledger = sum(json.loads(p.read_text(encoding='utf-8')).get('charge_usd', 0)
                 for p in (directory/'requests').glob('*.json'))
    assert abs(ledger-summary['estimated_charged_or_reserved_usd']) < 1e-9
    preserved = directory/'report_summary_generated.json'
    if not preserved.exists():
        preserved.write_bytes(original)
    # The generator's flag covers external calibration reuse; same-campaign
    # imports are already charged inside the local ledger too.
    summary['local_budget_includes_parent_charges'] = imported_charge > 0
    summary['imported_charge_accounting_note'] = (
        'Imported same-campaign request charges are included once in the local ledger; do not add parent copies again.')
    assert all(summary[k] == json.loads(original)[k] for k in json.loads(original)
               if k != 'local_budget_includes_parent_charges')
    save(summary_path, summary)
    note = Path(note_path).read_text(encoding='utf-8').strip()
    marker = '<!-- operational-recovery-context -->'
    markdown_path = directory/'REPORT.md'
    markdown = markdown_path.read_text(encoding='utf-8')
    if marker not in markdown:
        markdown_path.write_text(markdown+'\n\n'+marker+'\n## Operational recovery and charge accounting\n\n'+note+
            f'\n\nImported same-campaign charges (${imported_charge:.8f}) are already included in the ${ledger:.8f} local total. '
            'Copies in parent directories and the calibration checkpoint are not additional spending.\n', encoding='utf-8')
        from docx import Document
        document = Document(directory/'REPORT.docx')
        document.add_heading('Operational recovery and charge accounting', level=1)
        document.add_paragraph(note.replace('[Operational provenance](operational_io/)', 'Operational provenance (operational_io/)'))
        document.add_paragraph(f'Imported same-campaign charges (${imported_charge:.8f}) are already included in the '
                               f'${ledger:.8f} local total. Parent/checkpoint copies are not additional spending.')
        document.save(directory/'REPORT.docx')
        from openpyxl import load_workbook
        workbook = load_workbook(directory/'DETAILS.xlsx')
        sheet = workbook.create_sheet('Operational_context')
        sheet.append(['Context', 'Value'])
        sheet.append(['Recovery note', note])
        sheet.append(['Imported charges included once in local total', imported_charge])
        sheet.append(['Local campaign estimated total USD', ledger])
        for p in sorted((directory/'operational_io').glob('*.json')):
            sheet.append([p.name, p.read_text(encoding='utf-8')])
        workbook.save(directory/'DETAILS.xlsx')
    from openpyxl import load_workbook
    workbook = load_workbook(directory/'DETAILS.xlsx')
    for sheet in workbook:
        for row in sheet:
            for i, cell in enumerate(row[:-1]):
                if cell.value == 'local_budget_includes_parent_charges':
                    row[i+1].value = imported_charge > 0
    workbook.save(directory/'DETAILS.xlsx')
    save(directory/'report_annotation.json', dict(created_utc=datetime.now(timezone.utc).isoformat(),
         script=str(Path(__file__).resolve()), script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
         added_context=True, corrected_metadata_field='local_budget_includes_parent_charges',
         numeric_results_changed=False, raw_requests_or_scientific_sources_changed=False, api_calls=0,
         preserved_generated_summary=preserved.name))
    print('Operational report context added; numeric outcomes and raw evidence unchanged.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--note', required=True, type=Path)
    args = parser.parse_args()
    annotate(args.directory, args.note)
