"""Build a readable pilot brief from completed evidence; never constructs an API client."""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys

import pandas as pd
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from beergame_study.economy_report import analyze, markdown_table


def hashes(directory):
    return {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(directory.rglob('*')) if p.is_file()}


def build(directory):
    directory = Path(directory).resolve()
    before = hashes(directory)
    status = json.loads((directory / 'progress.json').read_text(encoding='utf-8'))
    if status['phase'] not in ('complete', 'stopped'):
        raise ValueError('Wait for completion or a stopped partial run before summarizing.')
    result = analyze(directory)
    if result['manifest']['protocol'] != 'student-high-economy-revision-v3':
        raise ValueError('This brief is only for the separately versioned economy protocol.')
    output = ROOT / 'reports' / directory.name
    if not (output / 'analysis_manifest.json').exists():
        raise ValueError('Generate the full saved-evidence report first.')
    document = Document()
    sections = []

    def paragraph(text):
        sections.append(text)
        document.add_paragraph(text)

    def heading(text, level=2):
        sections.append('#' * level + ' ' + text)
        document.add_heading(text, level=level-1)

    def table(frame):
        sections.append(markdown_table(frame))
        if frame.empty:
            document.add_paragraph('No observations available.')
            return
        rendered = frame.astype(object).where(pd.notna(frame), 'Not estimated')
        word = document.add_table(rows=1, cols=len(frame.columns))
        word.style = 'Light Shading Accent 1'
        for cell, name in zip(word.rows[0].cells, frame.columns):
            cell.text = name
        for row in rendered.itertuples(index=False, name=None):
            for cell, value in zip(word.add_row().cells, row):
                cell.text = f'{value:.6g}' if isinstance(value, float) else str(value)

    api = result['manifest']['config']['mode'] == 'api'
    heading('Paid LLM pilot: findings and next run' if api else 'Synthetic software example: no empirical findings', 1)
    rounds = result['rounds']
    model = rounds[rounds.phase == 'model'] if not rounds.empty else rounds
    complete = int(result['groups'].complete.sum()) if not result['groups'].empty else 0
    paragraph(f"Run {directory.name}; status {status['phase']}. Completed chains: {complete}; "
              f"{'Model-controlled' if api else 'Synthetic'} orders: {len(model)}. Calibration passed: {result['screen']['passed']}; "
              f"game computation separation passed: {result['main_manipulation_passed']}. "
              + ('Four fictional UB sophomore profiles were used, with the same GPT-6 Luna model and high effort in both conditions.'
                 if api else 'This run used deterministic synthetic responses; no model or human observations were collected.'))
    if status['phase'] == 'stopped':
        paragraph('Stop reason: ' + status.get('reason', 'See progress.json.'))

    heading('What happened to the hypotheses')
    paragraph('Game costs in the outcome tables are simulated holding/backlog costs, separate from API spending.')
    descriptions = {'H1_sensitivity': 'Pipeline sensitivity (revised H1)',
                    'H2_cost': 'Full scored cost (H2)', 'H3_bullwhip': 'Bullwhip (primary H3)'}
    summaries = []
    for row in result['hypotheses'].to_dict('records'):
        own = result['contrasts']
        own = own[own.hypothesis == row['hypothesis']] if not own.empty else own
        estimate = row['estimate']
        if pd.isna(estimate):
            direction = 'Not estimable'
        elif abs(estimate) < 1e-12:
            direction = 'No observed difference'
        else:
            expected_positive = row['predicted_short_minus_adequate'] == '>0'
            direction = 'Predicted direction' if (estimate > 0) == expected_positive else 'Opposite direction'
        summaries.append({'Outcome': descriptions[row['hypothesis']],
                          'Short': float(own.short.mean()) if not own.empty else None,
                          'Adequate': float(own.adequate.mean()) if not own.empty else None,
                          'Short - adequate': estimate, 'Observed direction': direction})
    table(pd.DataFrame(summaries))
    if not result['groups'].empty and result['groups'].bullwhip.notna().all() and (result['groups'].bullwhip < 1).all():
        paragraph('Every chain-average customer-referenced variance ratio is below 1 in this run. '
                  'A higher ratio between conditions is a relative difference in order variability; '
                  'it does not demonstrate chain-average amplification above customer-demand variance. '
                  'Individual local-demand ratios are reported separately in roles.csv.')
    paragraph('A single pair cannot estimate between-pair uncertainty; '
              'its observed direction is neither confirmation nor disproof. Weekly calls and the four roles are not independent students. '
              'The full report retains all eligibility checks and any available confidence intervals and adjusted p-values.')
    model_weeks = result['manifest']['config']['weeks'] - 4
    paragraph(f'Bullwhip is the mean of the four role order-variance/customer-demand-variance ratios over the {model_weeks} model weeks. '
              'Cost includes practice, model weeks and the 12-week automatic continuation. '
              'The revised H1 uses separate matched states with a 24-case pipeline difference; it does not identify the original dynamic beta.')

    heading('Completion, orders and actual computation')
    behavior = []
    if not model.empty:
        for condition, rows in model.groupby('condition'):
            revised = rows[rows.order_source == 'revision']
            starvation = rows[(rows.order == 0) & (rows.inventory == 0) &
                              (rows.supply_line == 0) & (rows.backlog > 0)]
            behavior.append({'Condition': condition, 'Orders': len(rows),
                             'Zero orders': int((rows.order == 0).sum()),
                             'Order range': f'{int(rows.order.min())} to {int(rows.order.max())}',
                             'Distinct quantities': int(rows.order.nunique()),
                             'Median total tokens': float(rows.total_output_tokens.median()),
                             'Changed by revision': int((revised.order != revised.initial_order).sum()),
                             'Zero despite empty stock/pipeline and backlog': len(starvation)})
    if behavior:
        behavior_frame = pd.DataFrame(behavior)
        table(behavior_frame[['Condition', 'Orders', 'Zero orders', 'Order range', 'Distinct quantities']])
        table(behavior_frame[['Condition', 'Median total tokens', 'Changed by revision',
                              'Zero despite empty stock/pipeline and backlog']])
    else:
        paragraph('No model weeks were committed.')
    ratio = result['actual_token_ratio']
    paragraph(f'Median short/adequate generated-token ratio: {ratio}; the prespecified maximum was 0.75. '
              'Both initial and revision tokens count. Initial ceiling: 16,384; additional revision ceiling: 8,192. '
              'These are ceilings, not observed usage. Temperature is omitted identically; the common 900-second timeout is technical infrastructure.')
    paragraph('Order zero remains a valid model choice. No preferred maximum, positive floor or artificial variation is imposed. '
              f"Fallback orders recorded: {status.get('fallback_decisions', 0)}. "
              'A completed game checks operation; it does not validate the profiles against actual students.')
    phase_costs = pd.DataFrame()
    if not rounds.empty:
        phase_costs = rounds.groupby(['condition', 'phase'], as_index=False).cost.sum()
        phase_costs.to_csv(output / 'phase_costs.csv', index=False)
        table(phase_costs)
        pivot = phase_costs.pivot(index='phase', columns='condition', values='cost')
        if {'short', 'adequate'}.issubset(pivot.columns):
            delta = pivot['short'] - pivot['adequate']
            total_delta = float(delta.sum())
            paragraph(f"Across recorded chains, the short-minus-adequate cost difference is "
                      f"{float(delta.get('model', 0)):.6g} during model weeks and "
                      f"{float(delta.get('automatic', 0)):.6g} during automatic continuation "
                      f"({total_delta:.6g} over all scored weeks).")
        paragraph('The automatic continuation repeats the last order and is not LLM behavior. '
                  'Compare its cost with the model-week cost before attributing the full-cost contrast to decisions throughout the game.')
    fits = result['beta_fits']
    if not fits.empty:
        usable = int(fits.usable.sum())
        complete_beta_chains = sum(len(g) == 4 and bool(g.usable.all())
                                  for _, g in fits.groupby(['block', 'condition']))
        paragraph(f'Original structural beta: {usable}/{len(fits)} role fits pass the existing usability checks; '
                  f'{complete_beta_chains} chains have four usable role fits. '
                  'An unusable coefficient is not evidence about supply-line neglect. The complete diagnostics remain in beta_fits.csv and the full report.')

    heading('What was spent')
    charges = {}
    for record in result['records']:
        stage = record['identity']['stage']
        charges[stage] = charges.get(stage, 0.0) + record.get('charge_usd', 0.0)
    table(pd.DataFrame([{'Stage': stage, 'Estimated USD': value} for stage, value in charges.items()]))
    total = status['prior_campaign_spend_usd'] + result['charge_usd']
    paragraph(f"This Luna run: ${result['charge_usd']:.8f}. Prior campaign spending: "
              f"${status['prior_campaign_spend_usd']:.8f}; combined campaign ledger: ${total:.8f}. "
              'These are conservative usage-based estimates including cache-write rates, not an invoice. '
              'They exclude taxes and unrelated account activity. The earlier $0.7198179 Mini calibration remains a separate failed-model screen.')
    paragraph('Hard guards: $2.20 for this run, $0.18 for its fresh calibration, $2.95 for all redesign runs together, '
              'and $10 for the original campaign including the earlier $5.19360675 pilot. '
              'Completed evidence and report generation are free to reuse; no new calls are made by this summary.')

    heading('Run it yourself in Jupyter')
    paragraph('Open notebooks/04_personality_deliberation_study.ipynb in the Beer Game (.venv) kernel. '
              'All paid flags default off. WATCH_RUNNING=True displays the background progress bar without buying calls. '
              'The Review results cell reads saved evidence without a key.')
    if api and result['screen']['passed']:
        paragraph("For a new prospective run, set MODE='api', RUN_CALIBRATION=False, RUN_GAMES=True, "
                  'choose PAIRS and WEEKS before running, and give REPLICATE_ID a new value. '
                  f"Set CALIBRATION_RUN=ROOT/'data/runs/{directory.name}' to reuse this passed calibration if sources are unchanged. "
                  'Enter the key only in the masked notebook prompt. Keep CAMPAIGN_ID unchanged; all prior spending still counts.')
    else:
        paragraph('This run cannot supply passed live calibration. Use a compatible passed API calibration or perform fresh calibration within the spending guards.')
    paragraph('Use a pair count divisible by four to balance profile-role assignments. '
              'The notebook power calculation requires a scientifically meaningful effect-size assumption; eight pairs is only the inference reporting floor. '
              'New replicate IDs create new API draws but the default seed repeats the same indexed demand paths. '
              'Do not pool incompatible horizons, models or protocols or count repeated demand paths as new demand scenarios.')

    heading('What this pilot can establish')
    paragraph('Passing completion and computation checks establishes a usable experiment on configured LLM participants. '
              'Extra revision combines additional computation with reconsideration; it does not isolate information overload, a pure single-call token cap, '
              'or human time pressure. No student data exist, and persona fidelity remains unvalidated. '
              'Replication may produce a positive, null, opposite or inconclusive result. Preserve every result and do not retune for the desired direction.')
    paragraph('Full analysis: REPORT.md / Economy_Study_Report.docx. Numerical tables: DETAILS.xlsx. '
              'Raw requests, exact prompts, usage and source snapshots: the source run under data/runs/. '
              'Design and primary literature: reports/revised_design_20261011/REPORT.md.')
    if (output / 'orders.png').exists():
        sections.append('![Order trajectories by role](orders.png)')
        document.add_picture(str(output / 'orders.png'), width=Inches(6))
    assert hashes(directory) == before, 'Summary changed raw evidence.'
    (output / 'FINDINGS.md').write_text('\n\n'.join(sections) + '\n', encoding='utf-8')
    document.save(output / 'Pilot_Findings.docx')
    provenance = dict(source_run=str(directory), raw_files_unchanged=True, raw_sha256=before,
                      api_calls=0, source_script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                      request_statuses=dict(Counter(r['status'] for r in result['records'])),
                      saved_evidence_audit=result['audit'])
    (output / 'findings_provenance.json').write_text(json.dumps(provenance, indent=2) + '\n', encoding='utf-8')
    print(output / 'FINDINGS.md')
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--run', type=Path, required=True)
    build(parser.parse_args().run)
