"""Export the dated literature/design review to Word without API calls."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
folder = ROOT/'reports/revised_design_20261011'
document = Document()
document.styles['Normal'].font.size = Pt(10)
for section in document.sections:
    section.left_margin = section.right_margin = Inches(.65)


def plain(text):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '').replace('`', '')


table = None
for line in (folder/'REPORT.md').read_text(encoding='utf-8').splitlines():
    if line.startswith('| '):
        fields = [plain(value.strip()) for value in line.strip('|').split('|')]
        if all(set(value) <= {'-', ':', ' '} for value in fields):
            continue
        if table is None:
            table = document.add_table(rows=1, cols=len(fields))
            table.style = 'Light Shading Accent 1'
            cells = table.rows[0].cells
        else:
            cells = table.add_row().cells
        for cell, value in zip(cells, fields):
            cell.text = value
        continue
    table = None
    if line.startswith('# '):
        document.add_heading(line[2:], 0)
    elif line.startswith('## '):
        document.add_heading(line[3:], 1)
    elif line:
        document.add_paragraph(plain(line))
document.save(folder/'Design_Review.docx')
print(folder/'Design_Review.docx')
