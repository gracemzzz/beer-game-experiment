"""Build the human-readable report for the completed October 5 pilot; no API calls."""
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from audit_completed_game import ROOT, ROLES

RUN = ROOT / 'data/runs/game_api_8236a9c840b29d36'
OUT = ROOT / 'reports/archive/20261005/pilot_02_game_api_8236a9c840b29d36'
audit = json.loads((OUT / 'audit.json').read_text())
manifest = json.loads((RUN / 'manifest.json').read_text())
frame = pd.read_csv(RUN / 'rounds.csv')
tail = pd.read_csv(RUN / 'continuation.csv')
agent = frame[frame.phase == 'agent']
groups = pd.read_csv(OUT / 'group_outcomes.csv').set_index('condition')
betas = pd.read_csv(OUT / 'beta_fits.csv')
roles = pd.read_csv(OUT / 'role_metrics.csv')
requests = pd.read_csv(OUT / 'request_index.csv')
manipulation = pd.read_csv(RUN / 'manipulation_check.csv').set_index('condition')
combined = pd.concat([frame, tail], ignore_index=True)
phase_costs = combined.groupby(['condition', 'phase']).agg(cost=('cost','sum'),
    inventory_units=('inventory','sum'), backlog_units=('backlog','sum')).reset_index()
phase_costs['holding_cost'] = .5 * phase_costs.inventory_units
phase_costs['backlog_cost'] = phase_costs.backlog_units
full_roles = combined.groupby(['condition','role']).agg(total_cost=('cost','sum'),
    inventory_units=('inventory','sum'), backlog_cost=('backlog','sum')).reset_index()
full_roles['holding_cost'] = .5 * full_roles.inventory_units
full_roles.to_csv(OUT / 'full_horizon_role_costs.csv', index=False)
phase_costs.to_csv(OUT / 'phase_costs.csv', index=False)
colors = {'short':'#c65335', 'adequate':'#267a93'}
labels = {'short':'Short: 96 tokens', 'adequate':'Adequate: 240 tokens'}
plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False})

# Shareable static figures, using observed results and plainly named windows.
fig, axes = plt.subplots(1, 2, figsize=(10.5, 3.8))
for i,c in enumerate(('short','adequate')):
    g = combined[combined.condition == c]
    holding = float((g.inventory * .5).sum()); backlog = float(g.backlog.sum())
    axes[0].bar(i, holding, color='#b5c3ca', label='Holding' if i == 0 else None)
    axes[0].bar(i, backlog, bottom=holding, color=colors[c], label='Backlog' if i == 0 else None)
    axes[0].text(i, holding+backlog+450, f'{holding+backlog:,.0f}', ha='center', fontweight='bold')
    axes[1].bar(i, manipulation.loc[c,'mean_preparation_tokens'], color=colors[c])
    axes[1].text(i, manipulation.loc[c,'mean_preparation_tokens']+5,
                 f"{manipulation.loc[c,'mean_preparation_tokens']:.1f}",ha='center')
    axes[1].plot([i-.35,i+.35],[manifest['config'][c+'_tokens']]*2, color='black',linestyle='--')
for ax in axes:
    ax.set_xticks([0,1], ['Short','Adequate'])
axes[0].set_ylim(0,32000); axes[0].set_ylabel('Game cost units (weeks 1–52)')
axes[0].set_title('Observed cost: one chain per condition')
axes[0].legend(frameon=False,fontsize=8)
axes[1].set_ylim(0,270); axes[1].set_ylabel('Mean memo output tokens')
axes[1].set_title('Actual use; dashed lines = allowed caps')
fig.tight_layout();fig.savefig(OUT / 'performance_and_tokens.png',dpi=180);plt.close(fig)

fig, axes = plt.subplots(2,2,figsize=(10.5,5.4),sharex=True)
for ax,role in zip(axes.flat,ROLES):
    for c in ('short','adequate'):
        g=frame[(frame.role==role)&(frame.condition==c)].sort_values('week')
        ax.plot(g.week,g.order,color=colors[c],label=labels[c],linewidth=1.5)
    ax.axvspan(1,4,color='#dddddd',alpha=.5)
    ax.set_title(role.title());ax.set_ylabel('Order units');ax.set_xlabel('Game week')
    ax.legend(fontsize=7,frameon=False)
fig.tight_layout();fig.savefig(OUT / 'orders_by_role.png',dpi=180);plt.close(fig)

fig,axes=plt.subplots(1,2,figsize=(10.5,3.5))
for c in ('short','adequate'):
    g=combined[(combined.condition==c)&(combined.role=='factory')].sort_values('week')
    axes[0].plot(g.week,g.backlog,color=colors[c],label=labels[c])
    gg=combined[combined.condition==c].groupby('week').cost.sum().cumsum()
    axes[1].plot(gg.index,gg.values,color=colors[c],label=labels[c])
for ax in axes:
    ax.axvline(40.5,color='gray',linestyle='--',linewidth=1)
    ax.set_xlabel('Week (41–52: automatic continuation)');ax.legend(fontsize=7,frameon=False)
axes[0].set_ylabel('Factory backlog units');axes[0].set_title('Short factory backlog persists')
axes[1].set_ylabel('Cumulative game cost units');axes[1].set_title('Cost gap exists before continuation')
fig.tight_layout();fig.savefig(OUT / 'backlog_and_cumulative_cost.png',dpi=180);plt.close(fig)

doc = Document()
section=doc.sections[0]
section.top_margin=Inches(.65);section.bottom_margin=Inches(.6)
section.left_margin=Inches(.7);section.right_margin=Inches(.7)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
style.paragraph_format.space_after=Pt(5)
for name in ('Title','Heading 1','Heading 2'):
    doc.styles[name].font.color.rgb=RGBColor.from_string('163B53')
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(12)
header=section.header.paragraphs[0]
header.text='BEER GAME  |  TOKEN-BUDGET PILOT 02'
header.style=doc.styles['Caption']
footer=section.footer.paragraphs[0]
footer.text='LLM simulation • one chain per condition • 5 October 2026                         Page '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def finding(label, text):
    para=doc.add_paragraph(style='List Number')
    para.add_run(label+' ').bold=True;para.add_run(text)

def table(headers, rows, widths=None, fontsize=9):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(t.rows[0].cells,headers):
        c.text=str(h)
        for run in c.paragraphs[0].runs:run.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,value in zip(cells,row):c.text=str(value)
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Inches(widths[j])
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs:run.font.size=Pt(fontsize)
    return t

def heading(text):doc.add_heading(text,level=1)
def page():doc.add_page_break()
def caption(text):p(text,'Caption')

doc.add_heading('Beer Game pilot: results and conclusions',0)
p('Completed 5 October 2026  |  Run game_api_8236a9c840b29d36')
finding('The run worked technically.', '2 chains × 4 role agents completed 288 decisions. All 8 startup checks passed first time. There were no failed requests, retries or replacement groups.')
finding('H2 has a clear descriptive result, not statistical confirmation.', 'The short-token chain cost 28,198 versus 5,247 game cost units: 5.37 times as much (+437.4%). One chain per condition cannot establish a general treatment effect.')
finding('H1 is unresolved.', 'The short factory’s beta is unreliable, so the planned short-chain beta and H1 comparison cannot be calculated. The other role comparisons have mixed directions.')
finding('The token restriction worked in this run.', 'Short memos hit the cap in 142/144 decisions (98.6%); adequate memos in 5/144 (3.5%). This measures visible preparation tokens, not human thinking time.')
finding('Unrealistic behavior remains.', 'The short factory repeatedly chose zero while stock and supply line were zero and backlog was large. These decisions cannot be explained simply by excess inventory. Both chains also overordered upstream.')
finding('Next step: refine feasibility before scaling.', 'Keep this pilot as evidence. Validate factory/backlog comprehension and beta measurement before a confirmatory H1/H2 study. This analysis made no paid API calls.')

heading('Results at a glance')
table(['Measure','Short (96 tokens)','Adequate (240 tokens)'],[
    ['Independent chains / role agents','1 / 4','1 / 4'],
    ['Agent orders (weeks 5–40)','144','144'],
    ['Total scored cost (weeks 1–52)','28,198','5,247'],
    ['Holding / backlog cost','546 / 27,652','3,827 / 1,420'],
    ['Zero orders / largest order','81/144 (56.3%) / 312','60/144 (41.7%) / 210'],
    ['Usable role beta fits','3 of 4','4 of 4'],
    ['End-of-week-52 stock / backlog¹','34 / 932','456 / 6']
])
p('API usage cost estimate: $0.277483; elapsed time: 586.5 seconds (9 min 47 sec). Game cost units are separate from API dollars. ¹ Stock/backlog totals sum all four roles; they are not unique end-customer orders.','Caption')

page();heading('1. What was tested, and what was verified')
table(['Item','Executed setting'],[
    ['Hypotheses from the reflection','H1: short allowance → lower supply-line weight β. H2: short allowance → higher team holding + backlog costs.'],
    ['Agents and information','Retailer, wholesaler, distributor, factory. All roles see shared physical inventories; their other operating data and histories remain role-specific.'],
    ['Decision period','40 game weeks: 4 fixed-order warmup weeks, then 36 agent-choice weeks. 12 additional automatic scored weeks.'],
    ['Physical settings','Initial stock 12 per role; flow 4. Order transmission 1 week; shipment 2 weeks; factory production 3 weeks. Customer demand 4 for weeks 1–4, then 8; future demand was not revealed.'],
    ['Objective','Minimize 0.5 × end-of-week inventory + 1 × end-of-week backlog across all roles/weeks. No revenue, purchasing charge or factory production charge.'],
    ['Model and limits','gpt-5.4-nano-2026-03-17; reasoning effort none in both conditions. Memo caps 96 / 240. Common final response cap 192. No advertised maximum order.'],
    ['Calibration','Passed fixed-cap calibration_api_6f7902f51127105b; earlier reference calibration retained. This game is a new feasibility pilot.']
],fontsize=9)
doc.add_heading('Independent evidence checks',level=2)
for text in [
    'Replayed all 320 role-week states and 96 continuation states from the recorded orders. Inventory, receipts, shipments, backlog, supply line and every cost matched.',
    'Matched all 288 final JSON orders and observed-state readbacks to the game rows; rebuilt all 576 decision request payloads, including each role’s history and memo.',
    'Verified 584 saved API responses and unique response IDs: 8 rules + 288 memos + 288 final replies. All final replies completed; 147 memo truncations were intended deadline outcomes.',
    'Confirmed the saved model, zero reported reasoning tokens, request limits, pricing arithmetic and unchanged source-code hashes. Raw evidence was preserved; this audit verifies saved records, not the provider’s billing dashboard.'
]:p(text,'List Bullet')
p('Display issue found: progress.json says decisions = 0 because that counter is never updated. The response files and game data contain all 288 decisions. No outcome is missing; the report uses the verified count. The experiment code and notebook were not changed during this analysis.','Caption')

heading('2. Did the token treatment actually differ?')
table(['Memo measure','Short','Adequate'],[
    ['Allowed output tokens','96','240'],
    ['Mean / median actual tokens','95.875 / 96','171.625 / 168'],
    ['Memos cut off at the cap','142/144 (98.6%)','5/144 (3.5%)'],
    ['Actual mean final-response tokens','39','39'],
    ['Hidden reasoning tokens reported','0','0']
])
p('The game meets the predeclared aggregate truncation checks (short ≥20%, adequate ≤5%, adequate uses more tokens). Adequate truncation was 1/72 early and 4/72 late: the late-half rate is 5.6%, a descriptive warning rather than a newly invented failure rule. No doubled-budget benchmark was run on these game trajectories, so adequate decision quality is not proven.')
p('The second API call sees the state and memo and can change the provisional order. This tests a two-stage preparation allowance; it is not an exact substitute for a 60–120 second human deadline.')

page();heading('3. Supply-chain performance and H2')
table(['Cost window','Short','Adequate','Difference'],[
    ['Weeks 1–4: fixed warmup','96','96','0'],
    ['Weeks 5–40: agent choices','16,047','2,607','13,440'],
    ['Weeks 41–52: automatic continuation','12,055','2,544','9,511'],
    ['Total scored','28,198','5,247','22,951']
])
p('H2 predicts a positive short-minus-adequate difference. The observed difference is +22,951, in that direction. It already exists before the automatic continuation. There is only one paired chain block, so there is no defensible estimate of between-chain variability and no confirmatory p-value or confidence interval.')
p('Backlog accounts for 98.1% of the short chain’s cost; holding stock accounts for 72.9% of the adequate chain’s cost. Factory costs explain 12,584 of the 22,951 gap (54.8%). The short chain underproduces at the factory while downstream roles accumulate/repeat orders; the adequate chain finishes with excess upstream stock.')
table(['Role: full 52-week cost','Short','Adequate'],[
    [role.title(),f"{full_roles[(full_roles.condition=='short')&(full_roles.role==role)].total_cost.iloc[0]:,.0f}",
                  f"{full_roles[(full_roles.condition=='adequate')&(full_roles.role==role)].total_cost.iloc[0]:,.0f}"] for role in ROLES
])
p('Both retailers shipped 304 units against 304 units of cumulative customer demand through week 40. That 100% cumulative fulfillment includes late deliveries. Retailer backlog occurred in 15 short-chain weeks and 13 adequate-chain weeks; it is not 100% on-time service.')
doc.add_picture(str(OUT / 'backlog_and_cumulative_cost.png'),width=Inches(6.9))
caption('Dashed line: end of agent-choice period. Automatic continuation continues scoring the existing state; it does not erase unfinished orders.')

heading('4. Beta and H1: why no conclusion is possible')
p('The fitted policy is order = max(0, forecast + α × (target − net inventory − β × supply line)). Within this model, β = 1 gives full relative supply-line weight; β = 0 gives none. H1 predicts lower β under the short allowance. Beta is a model-based behavioral estimate, not profit or a direct reading of the agent’s cognition.')
table(['Role','Short β','Adequate β','Interpretation'],[
    ['Retailer','0.005','≈0.000','Opposite direction, very small difference'],
    ['Wholesaler','0.591','0.370','Opposite direction'],
    ['Distributor','0.409','0.934','H1-predicted direction'],
    ['Factory','2.000 — invalid','≈0.000','Short fit fails; do not compare'],
    ['Planned four-role mean','Not available','0.326','H1 effect/p-value withheld']
],fontsize=9)
p('The short factory has α = 0.0126 (<0.02 minimum), β at its upper fitting bound of 2, and R² = 0.022. This is weak response/identification, not evidence of strong supply-line weighting. Seven fits pass the coded usability rules, but several sit on parameter boundaries and the adequate distributor explains only 12.6% of order variance. Usable does not mean substantively validated. Do not drop the factory or average only favorable roles to produce H1.')

page();heading('5. Large orders, zero orders and bullwhip')
p('There was no upper order-entry cap, and the emergency threshold of 100,000 was never reached. Orders of 312 and 210 are observed choices, not clipping at a configured maximum. Zero remains a genuine lower boundary: 141 of 288 orders were zero (49.0%). Some zeros follow excess inventory or a large existing pipeline; others reflect problematic decisions.')
table(['Evidence from saved states/replies','Meaning'],[
    ['Short distributor, week 28: backlog 312, stock 0, supply line 0; orders 312. Week 30: backlog still 312, supply line now 312; orders another 312.','Repeated replenishment despite an outstanding order is consistent with incomplete supply-line use. It is an example, not a causal estimate of H1.'],
    ['Short factory, weeks 29–38: ten consecutive zero orders with stock 0 and supply line 0. Backlog rises from 276 to 588.','Existing unmet orders are not resolved by zero new demand. Week 31’s memo emphasizes that production cannot reduce today’s backlog and overlooks its future benefit.'],
    ['Adequate factory, week 30: stock 110, backlog 0, demand 0; orders 0.','This zero has a recognizable inventory-control justification; zero by itself is not failure.'],
    ['Adequate factory, week 40: backlog 150, supply line 150; orders 210. Its memo treats the existing pipeline as irrelevant.','Supply-line misuse is not limited to the short condition. Passing the rules quiz and readbacks does not establish correct planning.']
],fontsize=9)
p('A state-based review found 13 zero decisions with positive backlog, zero inventory and zero supply line: 12 in the short chain (11 factory, 1 distributor) and 1 in the adequate chain (distributor). This is a post hoc diagnostic flag, not an exclusion rule. All are retained. Written memos are recorded explanations, not proof of hidden reasoning.')

doc.add_heading('Bullwhip: order variance divided by demand variance',level=2)
table(['Role','Short / customer','Adequate / customer','Short / local','Adequate / local'],[
    [role.title()]+[f"{roles[(roles.condition==c)&(roles.role==role)][metric].iloc[0]:,.2f}"
                   for metric,c in [('bullwhip_vs_customer','short'),('bullwhip_vs_customer','adequate'),
                                    ('bullwhip_vs_local','short'),('bullwhip_vs_local','adequate')]]
    for role in ROLES
],fontsize=9)
p('Ratios above 1 show amplification relative to the stated denominator. The adequate chain increases upstream; the short chain peaks at the distributor and falls at the factory because production is suppressed. Bullwhip need not increase at every stage. The customer-demand variance is only 1.477 (40-week sample variance), because demand steps once and then stays at 8. The very large ratios are descriptive and sensitive to that low-variance denominator; raw orders matter too.')

doc.add_heading('Offline economic checks: no additional API spending',level=2)
table(['Mechanical replay / reference policy','52-week cost'],[
    ['Every role orders its incoming demand (warmup unchanged)','156'],
    ['Every role orders zero after warmup','9,368'],
    ['Every role orders 40 after warmup','26,200'],
    ['Actual short-chain orders, exact replay','28,198'],
    ['Change only short factory week 29 from 0 to 276','17,845']
],fontsize=9)
p('The last check saves 10,353 cost units (36.7%) while holding every later agent order fixed and retaining the same automatic continuation. It shows that the recorded zero was costly under this replay; it is not an adaptive LLM rerun or the causal token effect. The demand-following benchmark is not claimed optimal and benefits from this simple demand path. Together these checks do not support the claim that the implemented cost function inevitably rewards zero or maximum orders.')

page();heading('6. What to do next')
finding('Keep the current result as a feasibility pilot.', 'API execution and the token contrast worked. Do not present this as confirmation of H1, H2 or human time pressure.')
finding('Resolve comprehension before buying more chains.', 'A versioned follow-up should make factory production mechanics and persistent backlog explicit, using the same factual clarification for both conditions. Check what happens with zero stock, zero pipeline and positive backlog, including states as large as those seen here. Passing numeric readbacks alone is insufficient.')
finding('Validate the H1 measurement prospectively.', 'Test the beta estimator on several known policies and representative simulated trajectories, including prolonged zero production and large backlogs. Decide in advance what counts as an interpretable fit. Do not loosen the fit checks or exclude this factory after seeing the result.')
finding('Recalibrate after changing decision prompts or the model.', 'Then run a fresh small pilot to check behavior and beta availability. Keep previous versions and failures. Do not choose prompts, budgets or models according to whether they produce significant H1/H2 results.')
finding('Scale only under a frozen protocol.', 'The reflection proposes 10 chains per condition (80 role agents). The code’s 8-pair minimum is only an inference floor, not a power calculation. One pair provides no between-chain variance estimate for planning power. Repeated weeks and the four roles are not independent treatment replicates.')
p('The current H2 comparison is a useful descriptive result. The H1 measurement problem cannot be solved merely by buying more of the same chains. No notebook settings, prompts, fit rules or paid-run IDs were changed for this report.')

heading('7. API cost and timing')
table(['Recorded measure','Value'],[
    ['API calls','584: 8 rules + 288 preparation + 288 final'],
    ['Input / output tokens','1,074,865 / 50,008'],
    ['Cached input / reasoning tokens','0 / 0'],
    ['Rates saved in this run','Input $0.20 per million; output $1.25 per million'],
    ['Calculated cost / local spending allowance','$0.277483 / $1.05'],
    ['Run elapsed time','586.5 seconds (9 min 47 sec)'],
    ['Mean request latency: memo / final','1.259 / 0.808 seconds'],
    ['Nominal pacing','One new request per second, shared across both chains']
])
p('Costs were recalculated from saved token usage using the run’s recorded rates; they are not a fresh price quote or an account balance. The earlier two calibrations, failed startup and this game total approximately $0.3720 in the four named runs only; this is not the cost of all historical experiments. This analysis cost $0 in API calls.')
p('If usage and pacing stayed similar, 8 fresh chain pairs would cost about $2.22 and take about 78 minutes; 10 pairs about $2.77 and 98 minutes. These are arithmetic extrapolations, exclude further calibration/retries and are not a recommendation to start the main study now.')

heading('8. Files, limitations and reproducibility')
p('Source: data/source/ThoughtsReflection_Paragraph.docx (H1/H2 and intended human design); docs/PROTOCOL.md (implemented adaptation); raw run folder data/runs/game_api_8236a9c840b29d36. The report makes no new claims about the background papers in the reflection.')
p('This is one model snapshot, a two-stage token task, one deterministic customer-demand path, shared inventory information, and a finite 12-week terminal policy. Total supply line is displayed, so the proposed human mechanism of aggregating pipeline stages is not isolated. No post-game questionnaires were collected. No model upgrade, prompt change or significance-driven rerun was performed.')
p('DETAILS.xlsx contains setup, cost windows, all role/beta/bullwhip metrics, token use, every game and continuation row, all 584 request summaries and texts, and the flagged zero-order states. audit.json records the replay checks and SHA-256 hashes of the preserved evidence. Run scripts/audit_completed_game.py to reproduce the read-only checks, then scripts/build_pilot_02_report.py to recreate this report.')

page();heading('Appendix: complete order trajectories')
doc.add_picture(str(OUT / 'orders_by_role.png'),width=Inches(6.9))
caption('Each panel uses its own vertical scale. Gray band: fixed-order warmup, weeks 1–4. All 36 agent-choice weeks remain included; no zero or large order was removed.')
doc.add_picture(str(OUT / 'performance_and_tokens.png'),width=Inches(6.9))
caption('One chain per condition. These charts show observed differences, not confidence intervals or estimated population effects.')

doc.core_properties.title='Beer Game token-budget pilot 02 — audited results'
doc.core_properties.subject='Conclusions first; H1/H2 status, operational diagnosis and full audit'
doc.core_properties.author='Research analysis'
doc.save(OUT / 'Pilot_Report.docx')

# Full data appendix in native spreadsheet tables rather than tiny tables in the report.
status = pd.DataFrame([dict(run_id=RUN.name, status='complete', groups=2, agents=8, verified_decisions=288,
    p_values='Withheld: one pair; H1 additionally lacks valid group beta',
    cost_usd=audit['estimated_api_cost_usd'], elapsed_s=audit['elapsed_s'])])
readme = pd.DataFrame([
    ('Scope','Real API pilot: one short chain and one adequate chain; not human data.'),
    ('Read first','Pilot_Report.docx begins with the conclusions and separates descriptive findings from inference.'),
    ('Units','Holding/backlog scores are game cost units. charge_usd is estimated API dollars.'),
    ('Beta','Short factory is invalid. Other usable fits still need boundary and R2 review.'),
    ('Window','Role metrics and bullwhip: weeks 1–40. Beta/zero/token metrics: 5–40. Full cost: 1–52.'),
    ('Legacy column','group_outcomes.agent_phase_cost includes warmup (weeks 1–40), despite its name.'),
    ('Progress counter',audit['progress_counter_note']),
    ('Zero review','Post hoc diagnostic, not exclusion. All original orders remain in analysis.'),
    ('Offline comparisons','Mechanical checks, not additional LLM observations, human results or optimized policies.'),
    ('Request files','request_index.file is relative to '+str(RUN/'requests')),
    ('No paid calls','Report generation only reads existing data and performs local calculations.'),
],columns=['Item','Explanation'])
cost_by_stage=requests.groupby(['condition','stage']).agg(requests=('file','size'),
    input_tokens=('input_tokens','sum'), output_tokens=('output_tokens','sum'),
    charge_usd=('charge_usd','sum'), mean_latency_s=('elapsed_s','mean')).reset_index()
token_by_period=agent.assign(period=np.where(agent.week<=22,'weeks 5–22','weeks 23–40')).groupby(
    ['condition','period']).agg(decisions=('order','size'),memo_tokens=('plan_tokens','mean'),
    truncated_fraction=('plan_truncated','mean')).reset_index()
tables={
    'Read me':readme, 'Status':status,
    'Configuration':pd.DataFrame([(k,str(v)) for k,v in manifest['config'].items()],columns=['Setting','Value']),
    'Group outcomes':groups.reset_index(), 'Cost by phase':phase_costs,
    'Full horizon role costs':full_roles, 'Hypothesis status':pd.read_csv(OUT/'hypotheses.csv'),
    'Beta diagnostics':betas,'Role and bullwhip':roles,
    'Token manipulation':manipulation.reset_index(),'Token by period':token_by_period,
    'API cost by stage':cost_by_stage,'Rules checks':pd.read_csv(RUN/'rules_validation.csv'),
    'Zero orders to review':pd.read_csv(OUT/'zero_order_review.csv'),
    'Offline comparisons':pd.read_csv(OUT/'offline_comparisons.csv'),
    'All game rows':frame,'Automatic continuation':tail,'All request texts':requests,
    'Evidence hashes':pd.DataFrame(audit['immutable_evidence_sha256'].items(),columns=['File','SHA256'])
}
with pd.ExcelWriter(OUT/'DETAILS.xlsx',engine='openpyxl') as writer:
    for name,data in tables.items():
        data.to_excel(writer,sheet_name=name,index=False)
        sheet=writer.sheets[name];sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cells in sheet.columns:
            letter=cells[0].column_letter;name_=str(cells[0].value)
            sheet.column_dimensions[letter].width=min(55,max(13,len(name_)+2,
                max((len(str(c.value)) for c in list(cells)[1:20] if c.value is not None),default=0)))
        for cell in sheet[1]:
            from openpyxl.styles import Font, PatternFill, Alignment
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163B53')
            cell.alignment=Alignment(wrap_text=True)
        sheet.row_dimensions[1].height=30
        if name=='Read me':
            sheet.column_dimensions['B'].width=110
            for row in sheet.iter_rows(min_row=2):
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
            for i in range(2,sheet.max_row+1):sheet.row_dimensions[i].height=32

(ROOT/'reports/LATEST_REPORT.txt').write_text(
    'Latest completed API pilot analysis\n'+str(OUT/'Pilot_Report.docx')+'\n'+str(OUT/'DETAILS.xlsx')+'\n'
    +'Raw run: '+str(RUN)+'\nNo paid API calls were made during analysis.\n')
assert all(hashlib.sha256((RUN/p).read_bytes()).hexdigest()==sha for p,sha in audit['immutable_evidence_sha256'].items())
print('Saved:',OUT/'Pilot_Report.docx')
print('Saved:',OUT/'DETAILS.xlsx')
print('Verified original evidence unchanged; no live API calls.')
