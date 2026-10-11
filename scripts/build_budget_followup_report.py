"""Report the budget campaign from saved evidence. No API calls."""
import hashlib
import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from run_budget_followup import ROOT, CAMPAIGN, DIRECTORY, totals

OUT=ROOT/'reports/budget_followup_20261005'

def build():
    OUT.mkdir(parents=True,exist_ok=True)
    campaign=json.loads(CAMPAIGN.read_text());spend=totals()
    screen=pd.read_csv(DIRECTORY/'summary.csv')
    comparison=pd.read_csv(DIRECTORY/'decisions.csv')
    cal=Path(campaign['calibration_dir'])
    rec=json.loads((cal/'calibration_recommendation.json').read_text())
    calibration=pd.read_csv(cal/'calibration_summary.csv')
    calibration_rows=pd.read_csv(cal/'calibration.csv')
    game=Path(campaign['game_dir']) if campaign.get('game_dir') else None
    done=game is not None and (game/'rounds.csv').exists() and json.loads((game/'progress.json').read_text())['phase']=='complete'
    source_paths=[DIRECTORY,cal]+([game] if game else [])
    recheck_dir=ROOT/'data/development/calibration_metadata_recheck_20261005'
    recheck=json.loads((recheck_dir/'result.json').read_text()) if (recheck_dir/'result.json').exists() else None
    if recheck:source_paths.append(recheck_dir)
    request_rows=[];hashes={}
    for folder in source_paths:
        for p in sorted((folder/'requests').glob('*.json')):
            r=json.loads(p.read_text());hashes[str(p.relative_to(ROOT))]=hashlib.sha256(p.read_bytes()).hexdigest()
            request_rows.append(dict(folder=folder.name,file=p.name,model=r['model'],stage=r['identity']['stage'],
                identity=json.dumps(r['identity']),status=r['status'],response_status=r.get('response_status'),
                input_tokens=r.get('usage',{}).get('input_tokens',0),output_tokens=r.get('usage',{}).get('output_tokens',0),
                reasoning_tokens=r.get('usage',{}).get('reasoning_tokens',0),charge_usd=r['charge_usd'],
                elapsed_s=r.get('elapsed_s'),text=r.get('text','')))
    request_table=pd.DataFrame(request_rows)
    assert abs(request_table.charge_usd.sum()-spend['new_spend_usd'])<1e-9
    software=pd.read_csv(ROOT/'reports/software_checks/beta_closed_loop_20261005/recovery.csv')
    doc=Document();sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
    doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    for name in ('Title','Heading 1','Heading 2'):doc.styles[name].font.color.rgb=RGBColor.from_string('163B53')
    sec.header.paragraphs[0].text='BEER GAME | BUDGET FOLLOW-UP | 5 OCTOBER 2026'
    foot=sec.footer.paragraphs[0];foot.text='LLM feasibility study; not human data.                                        Page '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    def p(text,style=None):return doc.add_paragraph(text,style)
    def finding(label,text):
        a=doc.add_paragraph(style='List Number');a.add_run(label+' ').bold=True;a.add_run(text)
    def heading(t):doc.add_heading(t,level=1)
    def table(headers,rows,size=9):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
        flag=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(flag)
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=str(v)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    for r in para.runs:r.font.size=Pt(size)
        return t
    doc.add_heading('Budget follow-up: conclusions first',0)
    p('Question: can better instructions or a different model improve a useful Beer Game test within approximately $3?')
    finding('The model comparison is complete.', '300 API calls compared 3 models × 2 prompt versions on the same 12 development states and factual quizzes. Cost/reservations: $0.2562.')
    finding('Choose GPT-5.6 Luna with clarified physical rules.', 'It passed both factual quizzes and all 24 final-response checks, with no zero production in the tested empty-pipeline factory shortages. Selection used feasibility and price, not H1/H2 results.')
    finding('A stronger model alone did not solve everything.', 'GPT-5.4 mini with the original instructions still produced one adequate-budget zero in a factory shortage. Clear rules helped nano and mini. Luna’s original 96-token allowance did not bind, so the new short candidate was fixed at 32 before fresh calibration.')
    finding('Fresh calibration '+('passed.' if rec['manipulation_screen_passed'] else 'failed.'),
        'Short/adequate caps are 32/240; the benchmark is 480. '+
        ('The new games use only this calibrated model/prompt.' if rec['manipulation_screen_passed'] else 'No full games were purchased; preserve the remaining budget.'))
    tables={'Campaign spending':pd.DataFrame(spend['stages']),'Model comparison':screen,
        'All comparison choices':comparison,'Calibration summary':calibration,'All calibration choices':calibration_rows,
        'Synthetic beta recovery':software,'Request ledger':request_table,
        'Evidence hashes':pd.DataFrame(hashes.items(),columns=['File','SHA256'])}
    if recheck:
        tables['Corrected calibration summary']=pd.read_csv(recheck_dir/'summary.csv')
        tables['Corrected calibration choices']=pd.read_csv(recheck_dir/'calibration_with_corrected_metadata.csv')
    if done:
        groups=pd.read_csv(game/'group_outcomes.csv');betas=pd.read_csv(game/'beta_fits.csv')
        frame=pd.read_csv(game/'rounds.csv');tail=pd.read_csv(game/'continuation.csv')
        agent=frame[frame.phase=='agent'];tokens=pd.read_csv(game/'manipulation_check.csv')
        rules=pd.read_csv(game/'rules_validation.csv');hyp=pd.read_csv(game/'hypotheses.csv')
        zeros=agent[(agent.order==0)&(agent.inventory==0)&(agent.supply_line==0)&(agent.backlog>0)]
        costs=groups.groupby('condition').total_scored_cost.mean()
        difference=float(costs['short']-costs['adequate'])
        usable=int(betas.usable.sum());total=len(betas)
        valid_groups=groups.group_beta.notna().sum()
        finding('Fresh games completed.', f"{len(groups)} chains, {total} role agents and {len(agent)} real decisions. Mean scored cost: short {costs['short']:,.1f}; adequate {costs['adequate']:,.1f}. "
            +('The cost difference is in the H2-predicted direction.' if difference>0 else 'The cost difference does not follow H2.'))
        finding('The previous zero/extreme-order pattern did not recur.',
            f"All {len(agent)} orders were positive and ranged from {int(agent.order.min())} to {int(agent.order.max())}. No order floor above zero or maximum quantity was imposed. This is improvement in this pilot, not proof of human realism.")
        finding('H1 is not identifiable from these game trajectories.', f'{usable}/{total} role beta fits pass the coded checks; {valid_groups}/{len(groups)} chains have a usable four-role mean. '
            +'The adequate agents exactly followed incoming demand, leaving no separate inventory-adjustment response from which to identify beta. More of this same trajectory would provide little new information about H1.')
        finding('H2 has a consistent pilot direction, not significance.',
            f"Mean short cost {costs['short']:,.0f} versus {costs['adequate']:,.0f}: {costs['short']/costs['adequate']:.2f} times as much. Both pairs show higher short cost, but two paired chain blocks do not establish statistical significance. No p-values are reported.")
        tables.update({'Group outcomes':groups,'Beta diagnostics':betas,'Token use in games':tokens,
            'Hypothesis status':hyp,'Role and bullwhip':pd.read_csv(game/'role_metrics.csv'),
            'Rules checks':rules,'Flagged zero orders':zeros,'All game rows':frame,'Automatic continuation':tail})
    finding('Spending remains within the agreed envelope.',
        f"New follow-up charges/reservations: ${spend['new_spend_usd']:.4f}. Including the four earlier current-study runs: ${spend['known_current_study_total_usd']:.4f}. "
        f"Approximately ${3-spend['known_current_study_total_usd']:.2f} of a $3 current-study allowance remains. Older archived experiments and the provider’s account balance are not included.")

    doc.add_page_break();heading('1. What changed, and how candidates were selected')
    p('The simulator, demand path and economic objective were unchanged. All four roles minimize 0.5 per inventory unit-week plus 1 per backlog unit-week. There is no production charge or order-entry maximum. New instructions explicitly distinguish new demand from existing backlog and state that factory production started now becomes stock three weeks later, without requiring an existing supply line.')
    p('The development screen used eight saved states from the previous pilot plus four new hypothetical states, two memo allowances (96 and 240), and a common 192-token final reply. Two repeated factual quizzes checked production timing and persistent backlog. There were no correctness retries or discarded candidate answers. These partly selected development states are not independent evidence of generalization.')
    table(['Model / prompt','Quiz passes','Valid finals','Adequate shortage zeros','Short shortage zeros','Eligible'],[
        [r.model.replace('gpt-','').replace('-2026-03-17','')+' / '+r.prompt,
         f'{r.quizzes_correct}/2',f'{r.valid_decisions}/24',r.adequate_factory_shortage_zeros,
         r.short_factory_shortage_zeros,'Yes' if r.eligible else 'No'] for r in screen.itertuples()],size=8)
    p('Eligibility required correct factual replies, valid final responses and positive production in all four adequate-allowance empty-pipeline factory shortages. Short-condition behavior was reported but did not select a model. Among eligible models, the predeclared rule chose the lowest documented input/output prices; clarified instructions were preferred within a model. No hypothesis estimate or p-value entered selection.')
    p('Luna’s 96-token development memos averaged 52.9 tokens with no truncation. That would not test a binding deadline. The 32-token candidate was chosen before the fresh calibration. The adequate allowance remained 240, and the doubled benchmark remained 480; the pass thresholds were not changed.')
    heading('2. Fresh calibration')
    table(['Screen','Mean memo tokens','Truncated','Mean order'],[
        [r.screen,f'{r.mean_tokens:.2f}',f'{r.truncated_fraction:.1%}',f'{r.mean_order:.2f}'] for r in calibration.itertuples()])
    p(f"Adequate versus 480-token benchmark: median absolute order difference {rec['median_absolute_adequate_benchmark_order_difference']:.1f} units; required ≤4. "
      'Other requirements: short truncation ≥20%, adequate truncation ≤5%, adequate actual tokens > short. All 24 states included ordinary and large-stock/backlog situations across all four roles. This was fixed-cap feasibility calibration, not human-response-time percentiles.')
    p('Both conditions use reasoning effort none. Visible memo output is restricted by the API; a second common 192-token stage supplies the order and numeric readbacks. This intervention does not isolate human cognitive effort or equate tokens to human seconds.')
    if recheck:
        p('Input-metadata audit: four hypothetical calibration states had positive stock and processed demand but a zero shipment field. '
          'The affected 12 decisions (four states × three caps) were rerun with consistent shipment values, after the games had started. '
          f"The corrected-data sensitivity check {'passed' if recheck['passed'] else 'failed'} the unchanged thresholds; median adequate/benchmark order difference was {recheck['median_absolute_adequate_benchmark_difference']:.1f}. "
          'This is a post hoc repair check, not a claim that the corrected calibration preceded the games. No live-game state or order was changed. '
          'The generator was then corrected for future calibrations; the exact executed code is archived. One stable-stock hypothetical development-screen state also used a simplified shipment field; it did not enter the factory-shortage eligibility criterion.')

    if done:
        doc.add_page_break();heading('3. Fresh game results')
        table(['Chain','Condition','Total cost','Game / continuation','Group β'],[
            [r.group_id,r.condition,f'{r.total_scored_cost:,.0f}',f'{r.agent_phase_cost:,.0f} / {r.continuation_cost:,.0f}',
             f'{r.group_beta:.3f}' if pd.notna(r.group_beta) else 'Unusable'] for r in groups.itertuples()])
        p('Game cost includes weeks 1–40 (including four fixed-order warmup weeks) plus 12 automatic continuation weeks. Each independent chain has four role agents; each agent made 36 choices. The chains in each block use the same external demand path. Earlier nano games were not pooled into this follow-up.')
        table(['Condition','Actual memo tokens','Truncated','Zero orders','Largest order'],[
            [r.condition,f'{r.mean_preparation_tokens:.2f}',f'{r.truncated_fraction:.1%}',f'{r.zero_rate:.1%}',int(r.max_order)]
            for r in tokens.itertuples()])
        p('Zeros can be appropriate with excess stock or substantial incoming replenishment. The workbook lists every concerning zero with positive backlog and no stock/pipeline. Costs, extreme orders and beta failures never triggered replacement or exclusion.')
        paired=groups.pivot(index='block',columns='condition',values='total_scored_cost')
        diff=paired['short']-paired['adequate']
        p('H2, paired short-minus-adequate cost: '+', '.join(f'block {i}: {v:+,.0f}' for i,v in diff.items())
          +f'. Mean difference {diff.mean():+,.1f}. '+('Both blocks have the predicted direction.' if (diff>0).all() else 'The block directions are not uniformly positive.'))
        if groups.group_beta.notna().all():
            bp=groups.pivot(index='block',columns='condition',values='group_beta');bd=bp['short']-bp['adequate']
            p('H1, paired short-minus-adequate beta: '+', '.join(f'block {i}: {v:+.3f}' for i,v in bd.items())
              +f'. Mean difference {bd.mean():+.3f}. H1 predicts a negative difference. These are descriptive estimates only.')
        else:p('H1’s planned paired estimate is unavailable because at least one chain lacks four usable role betas. Averaging only the convenient roles would change the planned outcome and is not done.')
        p('No inferential p-values: there are only two independent chain pairs. Hundreds of weekly decisions do not create hundreds of independent treatment replicates. Repeatedly trying models or prompts until a significant result appears would not establish the hypotheses.')
        doc.add_picture(str(game/'order_trajectories.png'),width=Inches(6.9))

        doc.add_page_break();heading('4. Beta diagnostics and remaining limitations')
        table(['Chain','Role','β','R²','Usable','Flag'],[
            [r.group_id,r.role,f'{r.beta:.3f}',f'{r.r2:.3f}','Yes' if r.usable else 'No',r.fit_flag]
            for r in betas.itertuples()],size=8)
        p('Beta is the fitted relative weight on outstanding replenishment in the structural ordering rule, not a direct measure of internal reasoning. The original fitting criteria were retained. Boundary estimates, poor fit or weak inventory response remain visible.')
        p('Why beta fails here: in the adequate chains, order equals incoming demand in every agent-choice week. The retailer orders 8 in every such week, so its order series has no variation. For other roles the fitted inventory-adjustment strength α is effectively zero. In order = forecast + α × (target − net inventory − β × supply line), beta has no observable influence when α = 0. Values near 1.504 from those fits are arbitrary/unidentified; they must not be interpreted as measured over-weighting. This differs from the previous nano factory’s poor operating behavior.')
        p(f"The free closed-loop check fitted 16 simulated role trajectories whose true beta was known. All passed the coded usability checks, but the largest absolute beta error was {software.absolute_beta_error.max():.3f} (true retailer beta 1.000, fitted 0.774). "
          'This shows why a numerical usability flag is not proof of unbiased measurement. Those synthetic tests are not GPT evidence.')
        p('Old-versus-new performance is descriptive: model, prompt and short-token allowance changed together. The development screen helps diagnose the changes, but no full-game causal effect of “model upgrade alone” can be claimed. The game still exposes total supply line directly, uses a deterministic demand step and finite terminal continuation, and collects no retrospective questionnaire.')
        doc.add_picture(str(game/'bullwhip.png'),width=Inches(6.6))
        p('Bullwhip is order variance / external-demand variance over the 40 game weeks, including warmup. Every adequate role has a local-demand variance ratio of exactly 1: it simply copies incoming demand. Its external-demand ratio rises upstream because the demand step arrives later within the fixed 40-week window; that is a window/delay effect, not extra local order amplification. The short chains show some genuine local amplification (up to 1.62).','Caption')

    doc.add_page_break();heading('5. Budget, execution and next decision')
    table(['Stage','Calls completed','Estimated charge/reservation'],[
        [r['stage'],r['done'],f"${r['cost_usd']:.4f}"] for r in spend['stages']])
    p(f"New total ${spend['new_spend_usd']:.4f}; known earlier current-study total $0.3720; combined ${spend['known_current_study_total_usd']:.4f}. "
      'The model screen had a $0.45 cap, calibration $0.12, games $0.90, and the metadata recheck $0.03. Their maximum sum plus earlier runs is below $3. No account balance was read. Pricing applies only to the saved model/rates; taxes and unrelated account usage are outside this ledger.')
    p('Luna requests include a conservative 25% input-cost allowance because the official documentation distinguishes cache-write charges. This makes the local spending figure an estimate/reservation, not an exact invoice. Retries retain their original request and cap; failed/unknown requests preserve their reservations. The SDK’s automatic retries were disabled.')
    if done:
        if not groups.group_beta.notna().all():
            finding('Do not spend the remainder on a large H1 experiment yet.', 'The beta endpoint still has missing chain estimates. More copies of the same run do not repair that measurement problem. Preserve the result and revise/validate the measurement plan prospectively before a new main study.')
            finding('A better next H1 measurement is controlled state variation.', 'Consider a separate matched-state task that varies inventory and outstanding replenishment independently while holding role, demand, backlog and token condition fixed. Define the relative sensitivity outcome and sample size before collecting data. This would be a new measurement task, not a retrospective replacement beta for these games. A richer, prospectively fixed demand path is another design option; do not choose it because it produces the preferred hypothesis direction.')
        else:
            finding('Use this as a pilot for a frozen main-study plan.', 'Four usable chain estimates are useful, but two paired blocks are too few to estimate power reliably. Define the effect worth detecting and a prospective sample size; do not infer readiness from a favorable sign alone.')
        finding('Retain H2 separately from H1.', 'Holding/backlog costs are observed even if beta is unusable. Report the direction, magnitude and between-block variation; do not claim statistical confirmation from this sample.')
    else:finding('Stop after the failed calibration.', 'The current token contrast did not pass the declared screen. No extra full games were purchased. A future candidate must be justified by feasibility and token usage rather than H1/H2 significance.')
    finding('Use notebook 02 to inspect the results.', '02_budget_model_comparison.ipynb opens in review-only mode. It shows spending, completion counts, model comparison, calibration and game tables. Completed stages are not purchased again.')
    heading('Evidence and sources')
    p('The supplied reflection remains the study source. The original physical simulator and beta estimator were used unchanged. The spreadsheet includes all candidate replies, calibration choices, full game data if available, pricing ledger and provenance hashes. No raw response or previous pilot was overwritten.')
    p('Independent replay matched all 576 final decisions, 640 game states and 192 continuation states. All 1,168 game requests succeeded without retries or replacement groups. The saved beta outputs were reproduced in the original Jupyter Python environment. A different SciPy environment changed tiny near-zero nuisance estimates, while all usable/unusable classifications agreed; unidentifiable beta values should not be overinterpreted. Exact executed follow-up scripts are in archive/budget_followup_20261005/.')
    p('Official OpenAI documentation, checked 5 October 2026:')
    for url in ['https://developers.openai.com/api/docs/models/gpt-5.4-mini',
                'https://developers.openai.com/api/docs/models/gpt-5.4-nano',
                'https://developers.openai.com/api/docs/models/gpt-5.6-luna']:
        p(url)
    p('Luna is documented under an undated model identifier; returned model metadata are retained with every response. Reproducibility across future model updates is therefore limited compared with a dated snapshot.')
    doc.save(OUT/'Budget_Followup_Report.docx')

    with pd.ExcelWriter(OUT/'DETAILS.xlsx',engine='openpyxl') as writer:
        for name,data in tables.items():
            data.to_excel(writer,sheet_name=name[:31],index=False)
            s=writer.sheets[name[:31]];s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
            for col in s.columns:
                s.column_dimensions[col[0].column_letter].width=min(55,max(13,len(str(col[0].value))+2,
                    max((len(str(c.value)) for c in list(col)[1:12] if c.value is not None),default=0)))
            for c in s[1]:
                c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='163B53');c.alignment=Alignment(wrap_text=True)
            s.row_dimensions[1].height=30
    save_data=dict(campaign=campaign,spending=spend,complete_games=done,request_files=len(request_table),
        source_hashes=hashes,analysis_api_calls=0)
    (OUT/'report_provenance.json').write_text(json.dumps(save_data,indent=2)+'\n')
    (ROOT/'reports/LATEST_REPORT.txt').write_text(str(OUT/'Budget_Followup_Report.docx')+'\n'+str(OUT/'DETAILS.xlsx')+'\n')
    print('Saved',OUT/'Budget_Followup_Report.docx')
    print(json.dumps(spend,indent=2))

if __name__=='__main__':build()
